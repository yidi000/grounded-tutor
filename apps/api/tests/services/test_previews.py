from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from pydantic import ValidationError
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from grounded_tutor.domain.answers import ImageLocator, PptxLocator, XlsxLocator
from grounded_tutor.domain.ingestion import ChunkSettings
from grounded_tutor.services.previews import PreviewError, preview_file, preview_text

FIXTURES = Path(__file__).parents[1] / "fixtures"


def _empty_docx() -> bytes:
    output = BytesIO()
    Document().save(output)
    return output.getvalue()


def _empty_pdf() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def _pdf_with_text(text: str) -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    stream = DecodedStreamObject()
    stream.set_data(f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode())
    page[NameObject("/Contents")] = stream
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {
                    NameObject("/F1"): DictionaryObject(
                        {
                            NameObject("/Type"): NameObject("/Font"),
                            NameObject("/Subtype"): NameObject("/Type1"),
                            NameObject("/BaseFont"): NameObject("/Helvetica"),
                        }
                    )
                }
            )
        }
    )
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def test_custom_separator_preview_is_deterministic() -> None:
    preview = preview_text(
        "Mean is an average.\n---\nMedian is the middle value.",
        ChunkSettings(
            training_type="chunk",
            setting_mode="custom",
            split_mode="char",
            chunk_size=1000,
            index_size=256,
            splitter="---",
        ),
    )

    assert [item.text for item in preview.items] == [
        "Mean is an average.",
        "Median is the middle value.",
    ]
    assert preview.authority == "estimated"


def test_rejects_chunk_size_outside_fastgpt_range() -> None:
    with pytest.raises(ValueError):
        ChunkSettings(
            training_type="chunk",
            setting_mode="custom",
            split_mode="size",
            chunk_size=50,
            index_size=32,
        )


def test_settings_serialize_with_only_the_public_fastgpt_aliases() -> None:
    settings = ChunkSettings(
        training_type="chunk",
        index_prefix_title=False,
        custom_pdf_parse=True,
        setting_mode="custom",
        split_mode="size",
        chunk_size=300,
        index_size=64,
        splitter="",
        qa_prompt="Keep definitions precise.",
    )

    assert settings.model_dump(by_alias=True) == {
        "trainingType": "chunk",
        "indexPrefixTitle": False,
        "customPdfParse": True,
        "chunkSettingMode": "custom",
        "chunkSplitMode": "size",
        "chunkSize": 300,
        "indexSize": 64,
        "chunkSplitter": "",
        "qaPrompt": "Keep definitions precise.",
    }


def test_settings_accept_exact_http_aliases() -> None:
    settings = ChunkSettings.model_validate(
        {
            "trainingType": "qa",
            "indexPrefixTitle": True,
            "customPdfParse": False,
            "chunkSettingMode": "auto",
            "chunkSplitMode": "paragraph",
            "chunkSize": 100,
            "indexSize": 32,
            "chunkSplitter": "",
            "qaPrompt": "Generate study questions.",
        }
    )

    assert settings.training_type == "qa"
    assert settings.qa_prompt == "Generate study questions."


@pytest.mark.parametrize(
    "payload",
    [
        {"trainingType": "other"},
        {"chunkSettingMode": "other"},
        {"chunkSplitMode": "other"},
        {"indexPrefixTitle": "true"},
        {"customPdfParse": 1},
        {"chunkSize": "1000"},
        {"indexSize": True},
        {"unknownSetting": "private"},
    ],
)
def test_settings_reject_unknown_literals_coercions_and_fields(payload: dict[str, object]) -> None:
    with pytest.raises(ValidationError):
        ChunkSettings.model_validate(payload)


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("chunkSize", 99),
        ("chunkSize", 3001),
        ("indexSize", 31),
        ("indexSize", 1001),
        ("chunkSplitter", "x" * 21),
        ("qaPrompt", "x" * 4001),
    ],
)
def test_settings_reject_invalid_validation_boundaries(field: str, value: object) -> None:
    payload: dict[str, object] = {
        "trainingType": "chunk",
        "chunkSize": 1000,
        "indexSize": 32,
    }
    payload[field] = value

    with pytest.raises(ValidationError):
        ChunkSettings.model_validate(payload)


def test_settings_accept_validation_boundaries_and_qa_small_chunk_size() -> None:
    minimum = ChunkSettings(chunk_size=100, index_size=32)
    maximum = ChunkSettings(chunk_size=3000, index_size=3000, splitter="x" * 20)
    qa = ChunkSettings(training_type="qa", chunk_size=50, index_size=32)

    assert minimum.chunk_size == 100
    assert maximum.index_size == 3000
    assert qa.chunk_size == 50


def test_custom_delimiter_requires_a_nonempty_splitter() -> None:
    with pytest.raises(ValidationError):
        ChunkSettings(setting_mode="custom", split_mode="char", splitter="")


def test_size_preview_is_ordered_and_lossless() -> None:
    text = "a" * 100 + "b" * 100 + "c" * 10
    preview = preview_text(
        text,
        ChunkSettings(setting_mode="custom", split_mode="size", chunk_size=100, index_size=32),
    )

    assert [item.text for item in preview.items] == ["a" * 100, "b" * 100, "c" * 10]
    assert "".join(item.text for item in preview.items) == text


def test_paragraph_preview_trims_empty_segments_and_splits_an_oversized_paragraph() -> None:
    preview = preview_text(
        f"  First paragraph.  \n\n\n{'x' * 110}\n\n Last paragraph. ",
        ChunkSettings(
            setting_mode="custom",
            split_mode="paragraph",
            chunk_size=100,
            index_size=32,
        ),
    )

    assert [item.text for item in preview.items] == [
        "First paragraph.",
        "x" * 100,
        "x" * 10,
        "Last paragraph.",
    ]


def test_automatic_preview_uses_paragraph_behavior() -> None:
    preview = preview_text("First.\n\nSecond.", ChunkSettings())

    assert [item.text for item in preview.items] == ["First.", "Second."]


def test_automatic_mode_ignores_inactive_custom_split_mode() -> None:
    preview = preview_text(
        f"{'a' * 150}\n\nSecond.",
        ChunkSettings(setting_mode="auto", split_mode="size", chunk_size=100, index_size=32),
    )

    assert [item.character_count for item in preview.items] == [150, 7]
    assert preview.items[-1].text == "Second."


def test_preview_bounds_items_and_excerpt_without_hiding_original_lengths() -> None:
    text = "\n\n".join(["x" * 600] * 60)
    preview = preview_text(
        text,
        ChunkSettings(
            setting_mode="custom",
            split_mode="paragraph",
            chunk_size=1000,
            index_size=32,
        ),
    )

    assert len(preview.items) == 50
    assert len(preview.items[0].text) == 500
    assert preview.items[0].character_count == 600
    assert preview.items[0].truncated is True
    assert [warning.code for warning in preview.warnings] == ["preview_truncated"]
    assert preview.character_count == len(text)


def test_qa_preview_returns_source_excerpts_and_never_local_question_answer_pairs() -> None:
    preview = preview_text(
        "Mean is an average.\n\nMedian is the middle value.",
        ChunkSettings(training_type="qa"),
    )

    assert [item.text for item in preview.items] == [
        "Mean is an average.",
        "Median is the middle value.",
    ]
    assert [warning.code for warning in preview.warnings] == [
        "qa_generated_after_processing"
    ]
    assert all(
        set(item.model_dump())
        == {"position", "text", "character_count", "truncated", "locator"}
        and item.locator is None
        for item in preview.items
    )


def test_txt_fixture_is_extracted_before_preview() -> None:
    content = (FIXTURES / "statistics.txt").read_bytes()

    preview = preview_file("statistics.TXT", content, ChunkSettings(), max_upload_bytes=10_000)

    assert preview.source_name == "statistics.TXT"
    assert [item.text for item in preview.items] == [
        "Mean is an average.",
        "Median is the middle value.",
    ]
    assert all(item.locator is None for item in preview.items)


def test_pptx_preview_preserves_one_record_per_nonempty_slide() -> None:
    preview = preview_file(
        "slides.pptx",
        (FIXTURES / "slides.pptx").read_bytes(),
        ChunkSettings(),
        max_upload_bytes=100_000,
    )

    assert [item.locator for item in preview.items] == [
        PptxLocator(slide=1, title="Chunking"),
        PptxLocator(slide=2, title="Grounding"),
    ]
    assert preview.items[0].text == "Chunking\nSplit source material into bounded records."
    assert preview.items[1].text == "Grounding\nCitations connect answers to evidence."


@pytest.mark.parametrize(
    ("content_kind", "expected_text"),
    [
        ("table", "Mean\nAverage"),
        ("nested_group", "Grouped evidence"),
    ],
)
def test_pptx_preview_extracts_table_and_nested_group_text(
    content_kind: str,
    expected_text: str,
) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    if content_kind == "table":
        table = slide.shapes.add_table(
            1, 2, Inches(1), Inches(1), Inches(6), Inches(1)
        ).table
        table.cell(0, 0).text = "Mean"
        table.cell(0, 1).text = "Average"
    else:
        outer = slide.shapes.add_group_shape()
        inner = outer.shapes.add_group_shape()
        inner.shapes.add_textbox(
            Inches(1), Inches(1), Inches(4), Inches(1)
        ).text = "Grouped evidence"
    output = BytesIO()
    presentation.save(output)

    preview = preview_file(
        "nested.pptx",
        output.getvalue(),
        ChunkSettings(),
        max_upload_bytes=100_000,
    )

    assert len(preview.items) == 1
    assert preview.items[0].locator == PptxLocator(slide=1)
    assert preview.items[0].text == expected_text


def test_office_qa_preview_keeps_actual_processing_warning() -> None:
    preview = preview_file(
        "slides.pptx",
        (FIXTURES / "slides.pptx").read_bytes(),
        ChunkSettings(training_type="qa"),
        max_upload_bytes=100_000,
    )

    assert [warning.code for warning in preview.warnings] == [
        "qa_generated_after_processing"
    ]


def test_xlsx_preview_preserves_sheet_and_used_cell_range() -> None:
    preview = preview_file(
        "scores.xlsx",
        (FIXTURES / "workbook.xlsx").read_bytes(),
        ChunkSettings(),
        max_upload_bytes=100_000,
    )

    assert len(preview.items) == 1
    assert preview.items[0].locator == XlsxLocator(
        sheet="Week 1", cell_range="A1:B3"
    )
    assert preview.items[0].text == "Topic | Score\nMean | 90\nMedian | 85"


def test_image_preview_requires_explicit_support_and_reports_metadata_only() -> None:
    content = (FIXTURES / "diagram.png").read_bytes()

    with pytest.raises(PreviewError) as caught:
        preview_file(
            "diagram.png", content, ChunkSettings(), max_upload_bytes=100_000
        )

    assert caught.value.code == "unsupported_file_type"

    preview = preview_file(
        "diagram.png",
        content,
        ChunkSettings(),
        max_upload_bytes=100_000,
        supports_image_files=True,
    )
    assert len(preview.items) == 1
    assert preview.items[0].locator == ImageLocator(filename="diagram.png")
    assert preview.items[0].text == "PNG image, 64 x 32 pixels"


def test_markdown_is_extracted_as_text() -> None:
    preview = preview_file(
        "notes.md",
        b"# Mean\n\nAn average.\n\n## Median\n\nThe middle value.",
        ChunkSettings(),
        max_upload_bytes=10_000,
    )

    assert [item.text for item in preview.items] == [
        "# Mean",
        "An average.",
        "## Median",
        "The middle value.",
    ]


def test_html_extracts_visible_text_without_script_or_style_content() -> None:
    content = (
        b"<html><head><style>.secret{}</style><script>privateToken</script></head>"
        b"<body><h1>Mean</h1><p>An average.</p></body></html>"
    )

    preview = preview_file(
        "notes.html", content, ChunkSettings(), max_upload_bytes=10_000
    )

    joined = "\n".join(item.text for item in preview.items)
    assert "Mean" in joined
    assert "An average." in joined
    assert "privateToken" not in joined
    assert ".secret" not in joined


def test_csv_extracts_rows_in_stable_column_order() -> None:
    preview = preview_file(
        "terms.csv",
        b'term,definition\nmean,"sum divided by count"\n',
        ChunkSettings(),
        max_upload_bytes=10_000,
    )

    assert preview.items[0].text == "term | definition\nmean | sum divided by count"


def test_docx_extracts_paragraphs_in_document_order() -> None:
    document = Document()
    document.add_paragraph("Mean is an average.")
    document.add_paragraph("Median is the middle value.")
    output = BytesIO()
    document.save(output)

    preview = preview_file(
        "statistics.docx", output.getvalue(), ChunkSettings(), max_upload_bytes=100_000
    )

    assert [item.text for item in preview.items] == [
        "Mean is an average.",
        "Median is the middle value.",
    ]


def test_docx_preserves_interleaved_paragraph_and_table_order() -> None:
    document = Document()
    document.add_paragraph("Before table.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Mean"
    table.cell(0, 1).text = "Average"
    document.add_paragraph("After table.")
    output = BytesIO()
    document.save(output)

    preview = preview_file(
        "statistics.docx", output.getvalue(), ChunkSettings(), max_upload_bytes=100_000
    )

    assert [item.text for item in preview.items] == [
        "Before table.",
        "Mean | Average",
        "After table.",
    ]


def test_pdf_extracts_page_text() -> None:
    preview = preview_file(
        "statistics.pdf", _pdf_with_text("Mean is an average."), ChunkSettings(), max_upload_bytes=100_000
    )

    assert preview.items[0].text == "Mean is an average."


def test_filename_is_reduced_to_its_basename() -> None:
    preview = preview_file(
        "../../private/statistics.txt",
        b"Mean is an average.",
        ChunkSettings(),
        max_upload_bytes=10_000,
    )

    assert preview.source_name == "statistics.txt"


def test_file_source_name_is_bounded_for_the_public_response() -> None:
    with pytest.raises(PreviewError) as caught:
        preview_file(
            f"{'x' * 252}.txt",
            b"Mean is an average.",
            ChunkSettings(),
            max_upload_bytes=10_000,
        )

    assert caught.value.code == "validation_error"


def test_size_limit_is_enforced_before_attempting_to_parse() -> None:
    with pytest.raises(PreviewError) as caught:
        preview_file(
            "broken.pdf",
            b"not-a-pdf",
            ChunkSettings(),
            max_upload_bytes=4,
        )

    assert caught.value.code == "file_too_large"


@pytest.mark.parametrize("filename", ["notes.exe", "notes", ".txt"])
def test_unsupported_or_missing_extension_is_rejected(filename: str) -> None:
    with pytest.raises(PreviewError) as caught:
        preview_file(filename, b"Mean", ChunkSettings(), max_upload_bytes=10_000)

    assert caught.value.code == "unsupported_file_type"


@pytest.mark.parametrize(
    ("filename", "content"),
    [
        ("empty.txt", b"  \n\t"),
        ("empty.md", b"\n\n"),
        ("empty.html", b"<html><script>onlyScript()</script></html>"),
        ("empty.csv", b",\n,\n"),
        ("empty.docx", _empty_docx()),
        ("empty.pdf", _empty_pdf()),
    ],
)
def test_empty_extracted_text_is_rejected(filename: str, content: bytes) -> None:
    with pytest.raises(PreviewError) as caught:
        preview_file(filename, content, ChunkSettings(), max_upload_bytes=100_000)

    assert caught.value.code == "empty_source"


def test_empty_pasted_text_is_rejected() -> None:
    with pytest.raises(PreviewError) as caught:
        preview_text(" \n\t ", ChunkSettings())

    assert caught.value.code == "empty_source"


def test_delimiter_only_text_is_rejected_instead_of_returning_an_empty_preview() -> None:
    settings = ChunkSettings(
        setting_mode="custom",
        split_mode="char",
        splitter="---",
    )

    with pytest.raises(PreviewError) as caught:
        preview_text(" --- \n--- ", settings)

    assert caught.value.code == "empty_source"


def test_encrypted_pdf_is_rejected_without_attempting_to_extract() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    writer.encrypt("not-public")
    output = BytesIO()
    writer.write(output)

    with pytest.raises(PreviewError) as caught:
        preview_file(
            "encrypted.pdf", output.getvalue(), ChunkSettings(), max_upload_bytes=100_000
        )

    assert caught.value.code == "unreadable_file"
    assert "not-public" not in f"{caught.value!r} {caught.value}"


@pytest.mark.parametrize(
    ("filename", "content"),
    [
        ("broken.pdf", b"%PDF parser-private-marker"),
        ("broken.docx", b"PK parser-private-marker"),
        ("broken.txt", b"\xff parser-private-marker"),
        ("broken.csv", b"\xff parser-private-marker"),
        ("broken.csv", b'"parser-private-marker'),
    ],
)
def test_malformed_or_unreadable_file_raises_stable_safe_error(
    filename: str, content: bytes
) -> None:
    with pytest.raises(PreviewError) as caught:
        preview_file(filename, content, ChunkSettings(), max_upload_bytes=100_000)

    public = f"{caught.value!r} {caught.value}"
    assert caught.value.code == "unreadable_file"
    assert "parser-private-marker" not in public


def test_automatic_preview_ignores_cached_custom_chunk_length() -> None:
    text = "a" * 1200
    cached_custom = ChunkSettings(
        setting_mode="auto", split_mode="char", splitter="a", chunk_size=100, index_size=32
    )
    preview = preview_text(text, cached_custom)
    assert preview == preview_text(text, ChunkSettings())
    assert [item.character_count for item in preview.items] == [1000, 200]
