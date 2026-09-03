import gc
import struct
import tracemalloc
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pydantic import ValidationError
from pypdf import PdfWriter

import grounded_tutor.services.previews as previews_module
from grounded_tutor.config import Settings
from grounded_tutor.domain.ingestion import ChunkSettings
from grounded_tutor.services.previews import PreviewError, PreviewService, preview_file

FIXTURES = Path(__file__).parents[1] / "fixtures"


class RecordingWorkspaceRepository:
    def __init__(self) -> None:
        self.get_calls = 0

    def get(self, workspace_id):
        self.get_calls += 1
        return object()


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("max_upload_bytes", 0),
        ("max_upload_bytes", -1),
        ("max_upload_bytes", True),
        ("max_preview_text_bytes", 0),
        ("max_preview_text_bytes", -1),
        ("max_preview_text_bytes", True),
        ("max_extracted_characters", 0),
        ("max_extracted_characters", -1),
        ("max_extracted_characters", True),
    ],
)
def test_preview_resource_limits_must_be_strictly_positive(
    field: str, value: int | bool
) -> None:
    with pytest.raises(ValidationError):
        Settings(**{field: value})


def test_multibyte_text_byte_budget_is_checked_before_workspace_lookup() -> None:
    repository = RecordingWorkspaceRepository()
    service = PreviewService(
        repository,
        max_upload_bytes=100,
        max_preview_text_bytes=4,
        max_extracted_characters=100,
    )

    with pytest.raises(PreviewError) as caught:
        service.from_text(
            "00000000-0000-0000-0000-000000000000",
            source_name="notes",
            text="界界",
            settings=ChunkSettings(),
        )

    assert caught.value.code == "text_too_large"
    assert repository.get_calls == 0


def test_text_character_budget_is_checked_before_workspace_lookup() -> None:
    repository = RecordingWorkspaceRepository()
    service = PreviewService(
        repository,
        max_upload_bytes=100,
        max_preview_text_bytes=100,
        max_extracted_characters=3,
    )

    with pytest.raises(PreviewError) as caught:
        service.from_text(
            "00000000-0000-0000-0000-000000000000",
            source_name="notes",
            text="four",
            settings=ChunkSettings(),
        )

    assert caught.value.code == "source_too_large"
    assert repository.get_calls == 0


def test_file_byte_budget_is_checked_before_workspace_lookup() -> None:
    repository = RecordingWorkspaceRepository()
    service = PreviewService(
        repository,
        max_upload_bytes=4,
        max_preview_text_bytes=100,
        max_extracted_characters=100,
    )

    with pytest.raises(PreviewError) as caught:
        service.from_file(
            "00000000-0000-0000-0000-000000000000",
            filename="notes.txt",
            content=b"12345",
            settings=ChunkSettings(),
        )

    assert caught.value.code == "file_too_large"
    assert repository.get_calls == 0


@pytest.mark.parametrize(
    ("filename", "content"),
    [
        ("large.txt", b"12345"),
        ("large.html", b"<p>12345</p>"),
        ("large.csv", b"1,2,3"),
    ],
)
def test_plain_extractors_share_the_extracted_character_budget(
    filename: str, content: bytes
) -> None:
    with pytest.raises(PreviewError) as caught:
        preview_file(
            filename,
            content,
            ChunkSettings(),
            max_upload_bytes=100,
            max_extracted_characters=4,
        )

    assert caught.value.code == "source_too_large"


def test_exact_20_mb_text_file_remains_within_the_supported_file_limit() -> None:
    preview = preview_file(
        "large-valid.txt",
        b"x" * 20_000_000,
        ChunkSettings(),
        max_upload_bytes=20_000_000,
    )

    assert preview.character_count == 20_000_000
    assert preview.items[0].text == "x" * 500


def test_small_highly_compressed_docx_is_rejected_before_python_docx_expands_it(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    document = Document()
    document.add_paragraph("A" * 3_000_000)
    output = BytesIO()
    document.save(output)
    assert len(output.getvalue()) < 100_000
    monkeypatch.setattr(
        previews_module,
        "Document",
        lambda stream: (_ for _ in ()).throw(
            AssertionError("python-docx must not open an unsafe archive")
        ),
    )

    with pytest.raises(PreviewError) as caught:
        preview_file(
            "compressed.docx",
            output.getvalue(),
            ChunkSettings(),
            max_upload_bytes=20_000_000,
        )

    assert caught.value.code == "unsafe_archive"


def test_docx_with_excessive_archive_members_is_rejected_before_zipfile_allocates(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    document = Document()
    document.add_paragraph("Mean is an average.")
    output = BytesIO()
    document.save(output)
    with ZipFile(output, "a", compression=ZIP_DEFLATED) as archive:
        for index in range(50_000):
            archive.writestr(f"customXml/extra-{index}.xml", "x")
    content = output.getvalue()
    del output
    gc.collect()
    monkeypatch.setattr(
        previews_module,
        "ZipFile",
        lambda stream: (_ for _ in ()).throw(
            AssertionError("ZipFile must not inspect an excessive central directory")
        ),
    )

    tracemalloc.start()
    try:
        with pytest.raises(PreviewError) as caught:
            preview_file(
                "many-members.docx",
                content,
                ChunkSettings(),
                max_upload_bytes=20_000_000,
            )
        _, peak_bytes = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()

    assert caught.value.code == "unsafe_archive"
    assert peak_bytes < 2_000_000


def test_docx_rejects_multidisk_eocd_before_zipfile(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    content = bytearray(_small_docx())
    eocd = content.rfind(b"PK\x05\x06")
    struct.pack_into("<H", content, eocd + 4, 1)
    _rejects_before_zipfile(bytes(content), monkeypatch)


def test_docx_rejects_forged_central_directory_bounds_before_zipfile(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    content = bytearray(_small_docx())
    eocd = content.rfind(b"PK\x05\x06")
    struct.pack_into("<I", content, eocd + 16, len(content) + 100)
    _rejects_before_zipfile(bytes(content), monkeypatch)


def test_docx_rejects_zip64_markers_before_zipfile(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    content = bytearray(_small_docx())
    eocd = content.rfind(b"PK\x05\x06")
    struct.pack_into("<H", content, eocd + 10, 0xFFFF)
    _rejects_before_zipfile(bytes(content), monkeypatch)


def test_docx_rejects_trailing_junk_with_fake_eocd_signature(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    content = _small_docx() + b"trailing-PK\x05\x06-junk"
    _rejects_before_zipfile(content, monkeypatch)


def test_docx_accepts_eocd_signature_inside_a_valid_zip_comment() -> None:
    output = BytesIO(_small_docx())
    with ZipFile(output, "a") as archive:
        archive.comment = b"comment contains PK\x05\x06 but is not an EOCD record"

    preview = preview_file(
        "commented.docx",
        output.getvalue(),
        ChunkSettings(),
        max_upload_bytes=20_000_000,
    )

    assert preview.items[0].text == "Mean is an average."


def test_docx_extracted_node_budget_is_deterministic() -> None:
    document = Document()
    for index in range(10_001):
        document.add_paragraph(f"Paragraph {index}")
    output = BytesIO()
    document.save(output)

    with pytest.raises(PreviewError) as caught:
        preview_file(
            "many-blocks.docx",
            output.getvalue(),
            ChunkSettings(),
            max_upload_bytes=20_000_000,
        )

    assert caught.value.code == "source_work_limit_exceeded"


def test_docx_table_work_budget_is_checked_before_reading_oversized_row(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    document = Document()
    document.add_paragraph("valid archive")
    output = BytesIO()
    document.save(output)

    class UnreadableCell:
        @property
        def text(self) -> str:
            raise AssertionError("cell text must not be read after work budget is exceeded")

    class OversizedRow:
        cells = (UnreadableCell(),) * 10_001

    class OversizedTable:
        rows = (OversizedRow(),)

    class FakeDocument:
        def iter_inner_content(self):
            yield OversizedTable()

    monkeypatch.setattr(previews_module, "Document", lambda stream: FakeDocument())
    monkeypatch.setattr(previews_module, "Table", OversizedTable)

    with pytest.raises(PreviewError) as caught:
        preview_file(
            "oversized-table.docx",
            output.getvalue(),
            ChunkSettings(),
            max_upload_bytes=20_000_000,
        )

    assert caught.value.code == "source_work_limit_exceeded"


@pytest.mark.parametrize(
    ("fixture", "filename", "parser_name"),
    [
        ("slides.pptx", "slides.pptx", "Presentation"),
        ("workbook.xlsx", "workbook.xlsx", "load_workbook"),
    ],
)
def test_office_formats_reject_corrupt_central_directory_before_parser(
    fixture: str,
    filename: str,
    parser_name: str,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    content = bytearray((FIXTURES / fixture).read_bytes())
    eocd = content.rfind(b"PK\x05\x06")
    struct.pack_into("<I", content, eocd + 16, len(content) + 100)
    monkeypatch.setattr(
        previews_module,
        parser_name,
        lambda *args, **kwargs: (_ for _ in ()).throw(
            AssertionError("unsafe package must not reach the Office parser")
        ),
    )

    with pytest.raises(PreviewError) as caught:
        preview_file(
            filename,
            bytes(content),
            ChunkSettings(),
            max_upload_bytes=20_000_000,
        )

    assert caught.value.code == "unsafe_archive"


@pytest.mark.parametrize(
    ("fixture", "filename"),
    [("slides.pptx", "slides.pptx"), ("workbook.xlsx", "workbook.xlsx")],
)
def test_office_formats_reject_decompression_bombs(
    fixture: str,
    filename: str,
) -> None:
    output = BytesIO((FIXTURES / fixture).read_bytes())
    with ZipFile(output, "a", compression=ZIP_DEFLATED) as archive:
        archive.writestr("custom/huge.bin", "A" * 3_000_000)

    with pytest.raises(PreviewError) as caught:
        preview_file(
            filename,
            output.getvalue(),
            ChunkSettings(),
            max_upload_bytes=20_000_000,
        )

    assert caught.value.code == "unsafe_archive"


def test_workbook_cell_work_cap_is_checked_before_iteration(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(previews_module, "MAX_XLSX_CELLS", 5)

    with pytest.raises(PreviewError) as caught:
        preview_file(
            "workbook.xlsx",
            (FIXTURES / "workbook.xlsx").read_bytes(),
            ChunkSettings(),
            max_upload_bytes=100_000,
        )

    assert caught.value.code == "source_work_limit_exceeded"


@pytest.mark.parametrize("content_kind", ["table", "nested_group"])
def test_pptx_nested_shapes_and_table_cells_count_against_work_cap(
    monkeypatch: pytest.MonkeyPatch,
    content_kind: str,
) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    if content_kind == "table":
        table = slide.shapes.add_table(
            1, 2, Inches(1), Inches(1), Inches(6), Inches(1)
        ).table
        table.cell(0, 0).text = "Mean"
        table.cell(0, 1).text = "Average"
    else:
        outer = slide.shapes.add_group_shape()
        inner = outer.shapes.add_group_shape()
        inner.shapes.add_textbox(
            Inches(1), Inches(1), Inches(4), Inches(1)
        ).text = "Grouped evidence"
    output = BytesIO()
    presentation.save(output)
    monkeypatch.setattr(previews_module, "MAX_PPTX_SHAPES", 2)

    with pytest.raises(PreviewError) as caught:
        preview_file(
            "nested.pptx",
            output.getvalue(),
            ChunkSettings(),
            max_upload_bytes=100_000,
        )

    assert caught.value.code == "source_work_limit_exceeded"


def test_pptx_shape_collection_cap_is_checked_before_position_reads(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class PositionedShape:
        top = 0
        left = 0

    class UnreadablePosition:
        @property
        def top(self):
            raise AssertionError("over-limit shape position must not be read")

        @property
        def left(self):
            raise AssertionError("over-limit shape position must not be read")

    class Shapes(list):
        title = None

    class Slide:
        shapes = Shapes([PositionedShape(), PositionedShape(), UnreadablePosition()])

    class FakePresentation:
        def __init__(self) -> None:
            self.slides = [Slide()]

    monkeypatch.setattr(previews_module, "Presentation", lambda stream: FakePresentation())
    monkeypatch.setattr(previews_module, "MAX_PPTX_SHAPES", 2)

    with pytest.raises(PreviewError) as caught:
        preview_file(
            "slides.pptx",
            (FIXTURES / "slides.pptx").read_bytes(),
            ChunkSettings(),
            max_upload_bytes=100_000,
        )

    assert caught.value.code == "source_work_limit_exceeded"


def test_pptx_table_cell_cap_is_checked_before_cell_text_reads(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class UnreadableCell:
        @property
        def text(self):
            raise AssertionError("over-limit cell text must not be read")

    class Row:
        cells = (UnreadableCell(), UnreadableCell())

    class Table:
        rows = (Row(),)

    class TableShape:
        top = 0
        left = 0
        shape_id = 1
        shape_type = 0
        has_text_frame = False
        has_table = True
        table = Table()

    class Shapes(list):
        title = None

    class Slide:
        shapes = Shapes([TableShape()])

    class FakePresentation:
        def __init__(self) -> None:
            self.slides = [Slide()]

    monkeypatch.setattr(previews_module, "Presentation", lambda stream: FakePresentation())
    monkeypatch.setattr(previews_module, "MAX_PPTX_SHAPES", 2)

    with pytest.raises(PreviewError) as caught:
        preview_file(
            "slides.pptx",
            (FIXTURES / "slides.pptx").read_bytes(),
            ChunkSettings(),
            max_upload_bytes=100_000,
        )

    assert caught.value.code == "source_work_limit_exceeded"


@pytest.mark.parametrize(
    "unsafe_member",
    ["xl/vbaProject.bin", "xl/externalLinks/externalLink1.xml"],
)
def test_workbook_rejects_macros_and_external_links(unsafe_member: str) -> None:
    output = BytesIO((FIXTURES / "workbook.xlsx").read_bytes())
    with ZipFile(output, "a", compression=ZIP_DEFLATED) as archive:
        archive.writestr(unsafe_member, b"not trusted")

    with pytest.raises(PreviewError) as caught:
        preview_file(
            "workbook.xlsx",
            output.getvalue(),
            ChunkSettings(),
            max_upload_bytes=100_000,
        )

    assert caught.value.code == "unreadable_file"


def test_enabled_image_preview_rejects_malformed_and_wrong_format() -> None:
    jpeg = BytesIO()
    Image.new("RGB", (4, 4), "white").save(jpeg, format="JPEG")

    for content in (b"not an image", jpeg.getvalue()):
        with pytest.raises(PreviewError) as caught:
            preview_file(
                "diagram.png",
                content,
                ChunkSettings(),
                max_upload_bytes=100_000,
                supports_image_files=True,
            )
        assert caught.value.code == "unreadable_file"


def test_image_decompression_limit_returns_resource_error(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(Image, "MAX_IMAGE_PIXELS", 100)

    with pytest.raises(PreviewError) as caught:
        preview_file(
            "diagram.png",
            (FIXTURES / "diagram.png").read_bytes(),
            ChunkSettings(),
            max_upload_bytes=100_000,
            supports_image_files=True,
        )

    assert caught.value.code == "source_work_limit_exceeded"


def test_pdf_page_budget_is_checked_before_text_extraction() -> None:
    writer = PdfWriter()
    for _ in range(501):
        writer.add_blank_page(width=100, height=100)
    output = BytesIO()
    writer.write(output)

    with pytest.raises(PreviewError) as caught:
        preview_file(
            "many-pages.pdf",
            output.getvalue(),
            ChunkSettings(),
            max_upload_bytes=20_000_000,
        )

    assert caught.value.code == "source_work_limit_exceeded"


def test_pdf_extracted_text_uses_the_shared_character_budget(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class FakePage:
        def extract_text(self) -> str:
            return "12345"

    class FakeReader:
        is_encrypted = False
        pages = (FakePage(),)

    monkeypatch.setattr(previews_module, "PdfReader", lambda stream: FakeReader())

    with pytest.raises(PreviewError) as caught:
        preview_file(
            "bounded.pdf",
            b"%PDF-valid-for-test-double",
            ChunkSettings(),
            max_upload_bytes=100,
            max_extracted_characters=4,
        )

    assert caught.value.code == "source_too_large"


def _small_docx() -> bytes:
    document = Document()
    document.add_paragraph("Mean is an average.")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _rejects_before_zipfile(content: bytes, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        previews_module,
        "ZipFile",
        lambda stream: (_ for _ in ()).throw(
            AssertionError("raw ZIP preflight must reject before ZipFile")
        ),
    )

    with pytest.raises(PreviewError) as caught:
        preview_file(
            "unsafe.docx",
            content,
            ChunkSettings(),
            max_upload_bytes=20_000_000,
        )

    assert caught.value.code == "unsafe_archive"
