from __future__ import annotations

import asyncio
from collections.abc import Callable
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from threading import Event, get_ident
from unittest.mock import AsyncMock
from uuid import UUID, uuid4

import pytest
from sqlalchemy import event as sqlalchemy_event
from sqlalchemy import select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session

from grounded_tutor.adapters.fakes import FakeCollection, FakeFastGPT
from grounded_tutor.adapters.fastgpt import (
    DatasetRef,
    ExternalServiceError,
    ProcessedChunk,
    RetrievedChunk,
)
from grounded_tutor.config import Settings
from grounded_tutor.db import create_database_engine
from grounded_tutor.domain.answers import GeneratedAnswer, GeneratedBlock, PptxLocator
from grounded_tutor.domain.ingestion import ChunkSettings
from grounded_tutor.domain.models import Base, Source, SourceStatus, SourceType, Workspace
from grounded_tutor.repositories.sources import (
    SourcePersistenceError,
    SourcePersistenceOutcome,
    SourceRepository,
)
from grounded_tutor.services.grounding import (
    MAX_LOCATOR_MARKER_CHARS,
    ReadyChunk,
    ground_generated_answer,
)
from grounded_tutor.services.previews import PreviewError, preview_file
from grounded_tutor.services.source_locks import WorkspaceIngestionBusyError, WorkspaceLockRegistry
from grounded_tutor.services.sources import (
    ExternalSourceServiceError,
    NormalizedMaterial,
    SourceLifecycleConflictError,
    SourceNotFoundError,
    SourceService,
    SourceWorkspaceNotFoundError,
    source_marker,
)


@pytest.fixture
def source_engine(tmp_path: Path):
    engine = create_database_engine(Settings(database_url=f"sqlite:///{tmp_path / 'sources.db'}"))
    Base.metadata.create_all(engine)
    try:
        yield engine
    finally:
        engine.dispose()


@pytest.fixture
def source_context(source_engine) -> tuple[Callable[[], Session], FakeFastGPT, Workspace, SourceService]:
    def session_factory() -> Session:
        return Session(source_engine, expire_on_commit=False)

    with session_factory() as session:
        workspace = Workspace(title="Statistics", dataset_id="dataset-statistics")
        session.add(workspace)
        session.commit()
        session.refresh(workspace)

    fake = FakeFastGPT()
    fake.datasets[workspace.dataset_id] = DatasetRef(workspace.dataset_id)
    session = session_factory()
    service = SourceService(
        SourceRepository(session),
        fake,
        WorkspaceLockRegistry(),
        max_upload_bytes=20_000_000,
        max_text_bytes=20_000_000,
    )
    try:
        yield session_factory, fake, workspace, service
    finally:
        session.close()


def _seed_review_replacement(
    source_engine,
) -> tuple[FakeFastGPT, Workspace, UUID, UUID]:
    lineage_id = uuid4()
    new_id = uuid4()
    with Session(source_engine, expire_on_commit=False) as session:
        workspace = Workspace(title="Statistics", dataset_id="dataset-statistics")
        session.add(workspace)
        session.flush()
        old = Source(
            id=lineage_id,
            workspace_id=workspace.id,
            name="notes",
            source_type=SourceType.TEXT,
            collection_id="collection-old",
            status=SourceStatus.READY,
            lineage_id=lineage_id,
            ingestion_config={},
        )
        new = Source(
            id=new_id,
            workspace_id=workspace.id,
            name="notes v2",
            source_type=SourceType.TEXT,
            collection_id="collection-new",
            status=SourceStatus.REVIEW,
            lineage_id=lineage_id,
            replaces_source_id=lineage_id,
            version=2,
            ingestion_config={},
        )
        session.add_all([old, new])
        session.commit()
    fake = FakeFastGPT()
    fake.datasets[workspace.dataset_id] = DatasetRef(workspace.dataset_id)
    fake.collections["collection-old"] = FakeCollection(
        dataset_id=workspace.dataset_id,
        name=f"{source_marker(lineage_id)}--notes",
        content="old",
        config={},
        chunks=[ProcessedChunk("chunk-old", "Old", "")],
        tags=(source_marker(lineage_id),),
        forbidden=False,
    )
    fake.collections["collection-new"] = FakeCollection(
        dataset_id=workspace.dataset_id,
        name=f"{source_marker(new_id)}--notes-v2",
        content="new",
        config={},
        chunks=[ProcessedChunk("chunk-new", "New", "")],
        tags=(source_marker(new_id),),
        forbidden=True,
    )
    return fake, workspace, lineage_id, new_id


def _seed_ready_source(source_engine) -> tuple[FakeFastGPT, Workspace, UUID]:
    source_id = uuid4()
    with Session(source_engine, expire_on_commit=False) as session:
        workspace = Workspace(title="Statistics", dataset_id="dataset-statistics")
        session.add(workspace)
        session.flush()
        source = Source(
            id=source_id,
            workspace_id=workspace.id,
            name="notes",
            source_type=SourceType.TEXT,
            collection_id="collection-ready",
            status=SourceStatus.READY,
            lineage_id=source_id,
            ingestion_config={},
        )
        session.add(source)
        session.commit()
    fake = FakeFastGPT()
    fake.datasets[workspace.dataset_id] = DatasetRef(workspace.dataset_id)
    fake.collections["collection-ready"] = FakeCollection(
        dataset_id=workspace.dataset_id,
        name=f"{source_marker(source_id)}--notes",
        content="ready",
        config={},
        chunks=[ProcessedChunk("chunk-ready", "Ready", "")],
        tags=(source_marker(source_id),),
        forbidden=False,
    )
    return fake, workspace, source_id


@pytest.mark.asyncio
async def test_reconciliation_survives_repeated_outer_cancellation(source_context) -> None:
    _, _, _, service = source_context
    started = asyncio.Event()
    release = asyncio.Event()
    completed = False

    async def reconcile() -> None:
        nonlocal completed
        started.set()
        await release.wait()
        completed = True

    outer = asyncio.create_task(
        service._run_reconciliation_to_completion(reconcile())
    )
    await asyncio.wait_for(started.wait(), timeout=1)
    outer.cancel()
    await asyncio.sleep(0)
    outer.cancel()
    await asyncio.sleep(0)
    waited_for_child = not outer.done()
    release.set()

    with pytest.raises(asyncio.CancelledError):
        await outer

    assert waited_for_child
    assert completed


@pytest.mark.asyncio
async def test_reconciliation_child_failure_preserves_outer_cancellation(
    source_context,
) -> None:
    _, _, _, service = source_context
    started = asyncio.Event()
    release = asyncio.Event()
    unretrieved: list[dict[str, object]] = []
    loop = asyncio.get_running_loop()
    previous_handler = loop.get_exception_handler()

    async def reconcile() -> None:
        started.set()
        await release.wait()
        raise RuntimeError("private child failure")

    loop.set_exception_handler(lambda _loop, context: unretrieved.append(context))
    try:
        outer = asyncio.create_task(
            service._run_reconciliation_to_completion(reconcile())
        )
        await asyncio.wait_for(started.wait(), timeout=1)
        outer.cancel()
        await asyncio.sleep(0)
        release.set()
        with pytest.raises(asyncio.CancelledError) as caught:
            await outer
        await asyncio.sleep(0)
    finally:
        loop.set_exception_handler(previous_handler)

    assert isinstance(caught.value.__cause__, RuntimeError)
    assert "private child failure" in str(caught.value.__cause__)
    assert not any(
        context.get("message") == "Task exception was never retrieved"
        for context in unretrieved
    )


@pytest.mark.asyncio
async def test_accept_flush_failure_restores_remote_review_state(source_engine) -> None:
    fake, workspace, old_id, new_id = _seed_review_replacement(source_engine)

    class MarkReadyFlushFails(Session):
        def flush(self, objects=None) -> None:
            if any(
                isinstance(instance, Source)
                and instance.id == new_id
                and instance.status is SourceStatus.READY
                for instance in self.dirty
            ):
                raise SQLAlchemyError("private mark-ready detail")
            super().flush(objects)

    session = MarkReadyFlushFails(bind=source_engine, expire_on_commit=False)
    service = SourceService(SourceRepository(session), fake, WorkspaceLockRegistry())
    try:
        with pytest.raises(SourcePersistenceError):
            await service.accept(workspace.id, new_id)
    finally:
        session.close()

    assert fake.collections["collection-new"].forbidden is True
    assert fake.collections["collection-old"].forbidden is False
    with Session(source_engine) as check:
        old = check.get(Source, old_id)
        new = check.get(Source, new_id)
    assert old is not None and old.superseded_at is None
    assert new is not None and new.status is SourceStatus.REVIEW


@pytest.mark.asyncio
async def test_accept_commit_then_error_uses_durable_ready_state(source_engine) -> None:
    fake, workspace, old_id, new_id = _seed_review_replacement(source_engine)

    class MarkReadyCommitThenError(Session):
        def __init__(self, *args, **kwargs) -> None:
            super().__init__(*args, **kwargs)
            self.failed = False
            sqlalchemy_event.listen(self, "after_commit", self._fail_once)

        def _fail_once(self, session) -> None:
            if not self.failed:
                self.failed = True
                raise SQLAlchemyError("private ambiguous commit detail")

    session = MarkReadyCommitThenError(bind=source_engine, expire_on_commit=False)
    service = SourceService(SourceRepository(session), fake, WorkspaceLockRegistry())
    try:
        result = await service.accept(workspace.id, new_id)
    finally:
        session.close()

    assert result.status is SourceStatus.READY
    assert fake.collections["collection-new"].forbidden is False
    assert fake.collections["collection-old"].forbidden is True
    with Session(source_engine) as check:
        old = check.get(Source, old_id)
        new = check.get(Source, new_id)
    assert old is not None and old.superseded_at is not None
    assert new is not None and new.status is SourceStatus.READY


@pytest.mark.asyncio
async def test_accept_confirms_remote_enable_applied_before_timeout(source_engine) -> None:
    fake, workspace, _, new_id = _seed_review_replacement(source_engine)
    original_set = fake.set_collection_forbidden
    timed_out = False

    async def apply_then_timeout(collection_id: str, forbidden: bool) -> None:
        nonlocal timed_out
        await original_set(collection_id, forbidden)
        if collection_id == "collection-new" and not forbidden and not timed_out:
            timed_out = True
            raise ExternalServiceError(
                service="fastgpt",
                category="timeout",
                safe_message="FastGPT request failed.",
            )

    fake.set_collection_forbidden = apply_then_timeout
    with Session(source_engine, expire_on_commit=False) as session:
        service = SourceService(SourceRepository(session), fake, WorkspaceLockRegistry())
        result = await service.accept(workspace.id, new_id)

    assert result.status is SourceStatus.READY
    assert fake.collections["collection-new"].forbidden is False
    assert fake.collections["collection-old"].forbidden is True


@pytest.mark.asyncio
async def test_accept_retries_failed_compensation_until_remote_is_confirmed(
    source_engine,
) -> None:
    fake, workspace, _, new_id = _seed_review_replacement(source_engine)
    original_set = fake.set_collection_forbidden
    compensation_attempts = 0

    async def fail_old_and_first_compensation(
        collection_id: str, forbidden: bool
    ) -> None:
        nonlocal compensation_attempts
        if collection_id == "collection-old" and forbidden:
            raise RuntimeError("private old disable failure")
        if collection_id == "collection-new" and forbidden:
            compensation_attempts += 1
            if compensation_attempts == 1:
                raise RuntimeError("private first compensation failure")
        await original_set(collection_id, forbidden)

    fake.set_collection_forbidden = fail_old_and_first_compensation
    with Session(source_engine, expire_on_commit=False) as session:
        service = SourceService(SourceRepository(session), fake, WorkspaceLockRegistry())
        with pytest.raises(ExternalSourceServiceError) as caught:
            await service.accept(workspace.id, new_id)

    assert "private" not in str(caught.value)
    assert compensation_attempts == 2
    assert fake.collections["collection-new"].forbidden is True
    assert fake.collections["collection-old"].forbidden is False


@pytest.mark.asyncio
async def test_accept_cancellation_restores_remote_review_state(source_engine) -> None:
    fake, workspace, _, new_id = _seed_review_replacement(source_engine)
    original_set = fake.set_collection_forbidden
    enabled = asyncio.Event()

    async def apply_enable_then_block(collection_id: str, forbidden: bool) -> None:
        await original_set(collection_id, forbidden)
        if collection_id == "collection-new" and not forbidden:
            enabled.set()
            await asyncio.Event().wait()

    fake.set_collection_forbidden = apply_enable_then_block
    with Session(source_engine, expire_on_commit=False) as session:
        service = SourceService(SourceRepository(session), fake, WorkspaceLockRegistry())
        task = asyncio.create_task(service.accept(workspace.id, new_id))
        await asyncio.wait_for(enabled.wait(), timeout=1)
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task

    assert fake.collections["collection-new"].forbidden is True
    assert fake.collections["collection-old"].forbidden is False
    with Session(source_engine) as check:
        new = check.get(Source, new_id)
    assert new is not None and new.status is SourceStatus.REVIEW


@pytest.mark.asyncio
async def test_accept_persistence_compensation_survives_cancellation(
    source_engine,
) -> None:
    fake, workspace, old_id, new_id = _seed_review_replacement(source_engine)
    compensation_started = asyncio.Event()
    release_compensation = asyncio.Event()
    original_set = fake.set_collection_forbidden

    class MarkReadyFlushFails(Session):
        def flush(self, objects=None) -> None:
            if any(
                isinstance(instance, Source)
                and instance.id == new_id
                and instance.status is SourceStatus.READY
                for instance in self.dirty
            ):
                raise SQLAlchemyError("private mark-ready detail")
            super().flush(objects)

    async def block_compensation(collection_id: str, forbidden: bool) -> None:
        if collection_id == "collection-new" and forbidden:
            compensation_started.set()
            await release_compensation.wait()
        await original_set(collection_id, forbidden)

    fake.set_collection_forbidden = block_compensation
    session = MarkReadyFlushFails(bind=source_engine, expire_on_commit=False)
    service = SourceService(SourceRepository(session), fake, WorkspaceLockRegistry())
    try:
        task = asyncio.create_task(service.accept(workspace.id, new_id))
        await asyncio.wait_for(compensation_started.wait(), timeout=1)
        task.cancel()
        await asyncio.sleep(0)
        waited_for_compensation = not task.done()
        release_compensation.set()
        with pytest.raises(asyncio.CancelledError):
            await task
    finally:
        session.close()

    assert waited_for_compensation
    assert fake.collections["collection-new"].forbidden is True
    assert fake.collections["collection-old"].forbidden is False
    with Session(source_engine) as check:
        old = check.get(Source, old_id)
        new = check.get(Source, new_id)
    assert old is not None and old.superseded_at is None
    assert new is not None and new.status is SourceStatus.REVIEW


@pytest.mark.asyncio
async def test_accept_remote_failure_compensation_survives_cancellation(
    source_engine,
) -> None:
    fake, workspace, old_id, new_id = _seed_review_replacement(source_engine)
    compensation_started = asyncio.Event()
    release_compensation = asyncio.Event()
    original_set = fake.set_collection_forbidden

    async def fail_old_and_block_compensation(
        collection_id: str, forbidden: bool
    ) -> None:
        if collection_id == "collection-old" and forbidden:
            raise RuntimeError("private old disable failure")
        if collection_id == "collection-new" and forbidden:
            compensation_started.set()
            await release_compensation.wait()
        await original_set(collection_id, forbidden)

    fake.set_collection_forbidden = fail_old_and_block_compensation
    with Session(source_engine, expire_on_commit=False) as session:
        service = SourceService(SourceRepository(session), fake, WorkspaceLockRegistry())
        task = asyncio.create_task(service.accept(workspace.id, new_id))
        await asyncio.wait_for(compensation_started.wait(), timeout=1)
        task.cancel()
        await asyncio.sleep(0)
        waited_for_compensation = not task.done()
        release_compensation.set()
        with pytest.raises(asyncio.CancelledError):
            await task

    assert waited_for_compensation
    assert fake.collections["collection-new"].forbidden is True
    assert fake.collections["collection-old"].forbidden is False
    with Session(source_engine) as check:
        old = check.get(Source, old_id)
        new = check.get(Source, new_id)
    assert old is not None and old.superseded_at is None
    assert new is not None and new.status is SourceStatus.REVIEW


@pytest.mark.asyncio
async def test_accept_durable_reread_failure_restores_conservative_remote_state(
    source_engine,
) -> None:
    fake, workspace, old_id, new_id = _seed_review_replacement(source_engine)

    class MarkReadyCommitThenError(Session):
        def __init__(self, *args, **kwargs) -> None:
            super().__init__(*args, **kwargs)
            sqlalchemy_event.listen(self, "after_commit", self._fail)

        def _fail(self, session) -> None:
            raise SQLAlchemyError("private ambiguous commit detail")

    class HistoricalReadFails(SourceRepository):
        def get_historical_for_workspace(self, workspace_id, source_id):
            raise SourcePersistenceError(
                outcome=SourcePersistenceOutcome.DEFINITELY_UNCOMMITTED
            )

    session = MarkReadyCommitThenError(bind=source_engine, expire_on_commit=False)
    service = SourceService(HistoricalReadFails(session), fake, WorkspaceLockRegistry())
    try:
        with pytest.raises(SourcePersistenceError) as caught:
            await service.accept(workspace.id, new_id)
    finally:
        session.close()

    assert "private" not in str(caught.value)
    assert fake.collections["collection-new"].forbidden is True
    assert fake.collections["collection-old"].forbidden is False
    with Session(source_engine) as check:
        old = check.get(Source, old_id)
        new = check.get(Source, new_id)
    assert old is not None and old.superseded_at is not None
    assert new is not None and new.status is SourceStatus.READY


@pytest.mark.asyncio
async def test_delete_flush_failure_restores_remote_enabled_state(source_engine) -> None:
    fake, workspace, source_id = _seed_ready_source(source_engine)

    class MarkDeletedFlushFails(Session):
        def flush(self, objects=None) -> None:
            if any(
                isinstance(instance, Source) and instance.deleted_at is not None
                for instance in self.dirty
            ):
                raise SQLAlchemyError("private mark-deleted detail")
            super().flush(objects)

    session = MarkDeletedFlushFails(bind=source_engine, expire_on_commit=False)
    service = SourceService(SourceRepository(session), fake, WorkspaceLockRegistry())
    try:
        with pytest.raises(SourcePersistenceError):
            await service.delete(workspace.id, source_id)
    finally:
        session.close()

    assert fake.collections["collection-ready"].forbidden is False
    with Session(source_engine) as check:
        stored = check.get(Source, source_id)
    assert stored is not None and stored.deleted_at is None


@pytest.mark.asyncio
async def test_delete_commit_then_error_uses_durable_deleted_state(source_engine) -> None:
    fake, workspace, source_id = _seed_ready_source(source_engine)

    class MarkDeletedCommitThenError(Session):
        def __init__(self, *args, **kwargs) -> None:
            super().__init__(*args, **kwargs)
            self.failed = False
            sqlalchemy_event.listen(self, "after_commit", self._fail_once)

        def _fail_once(self, session) -> None:
            if not self.failed:
                self.failed = True
                raise SQLAlchemyError("private ambiguous delete detail")

    session = MarkDeletedCommitThenError(bind=source_engine, expire_on_commit=False)
    service = SourceService(SourceRepository(session), fake, WorkspaceLockRegistry())
    try:
        await service.delete(workspace.id, source_id)
    finally:
        session.close()

    assert fake.collections["collection-ready"].forbidden is True
    with Session(source_engine) as check:
        stored = check.get(Source, source_id)
    assert stored is not None and stored.deleted_at is not None


@pytest.mark.asyncio
async def test_delete_restore_failure_returns_safe_error_without_local_delete(
    source_engine,
) -> None:
    fake, workspace, source_id = _seed_ready_source(source_engine)

    class MarkDeletedFlushFails(Session):
        def flush(self, objects=None) -> None:
            if any(
                isinstance(instance, Source) and instance.deleted_at is not None
                for instance in self.dirty
            ):
                raise SQLAlchemyError("private mark-deleted detail")
            super().flush(objects)

    original_set = fake.set_collection_forbidden

    async def fail_restore(collection_id: str, forbidden: bool) -> None:
        if not forbidden:
            raise RuntimeError("private restore detail")
        await original_set(collection_id, forbidden)

    fake.set_collection_forbidden = fail_restore
    session = MarkDeletedFlushFails(bind=source_engine, expire_on_commit=False)
    service = SourceService(SourceRepository(session), fake, WorkspaceLockRegistry())
    try:
        with pytest.raises(ExternalSourceServiceError) as caught:
            await service.delete(workspace.id, source_id)
    finally:
        session.close()

    assert "private" not in str(caught.value)
    assert fake.collections["collection-ready"].forbidden is True
    with Session(source_engine) as check:
        stored = check.get(Source, source_id)
    assert stored is not None and stored.deleted_at is None


@pytest.mark.asyncio
async def test_new_collection_is_disabled_until_acceptance(source_context) -> None:
    _, fake, workspace, service = source_context

    result = await service.ingest_text(
        workspace_id=workspace.id,
        name="week-1",
        text="Mean is an average.",
        settings=ChunkSettings(),
    )

    assert result.source.status is SourceStatus.REVIEW
    assert result.source.source_type is SourceType.TEXT
    assert fake.set_collection_forbidden_calls == [("collection-1", True)]
    assert fake.list_collection_data_calls == [("collection-1", 30)]
    assert result.processed_preview.items[0].q == "Mean is an average."
    assert result.processed_preview.authority == "actual"
    marker = source_marker(result.source.id)
    assert fake.call_history == [
        ("create_text_collection", "dataset-statistics", f"{marker}--week-1"),
        ("set_collection_forbidden", "collection-1", True),
        ("list_collections", "dataset-statistics", 0, 30, marker),
        ("list_collection_data", "collection-1", 30),
    ]


def test_ready_collection_ids_excludes_noncurrent_and_other_workspace_sources(
    source_context,
) -> None:
    session_factory, _, workspace, _ = source_context
    with session_factory() as session:
        other = Workspace(title="Other", dataset_id="dataset-other")
        session.add(other)
        session.flush()
        active = Source(
            workspace_id=workspace.id,
            name="active",
            source_type=SourceType.TEXT,
            collection_id="active",
            status=SourceStatus.READY,
            ingestion_config={},
        )
        session.add_all(
            [
                active,
                Source(
                    workspace_id=workspace.id,
                    name="review",
                    source_type=SourceType.TEXT,
                    collection_id="review",
                    status=SourceStatus.REVIEW,
                    ingestion_config={},
                ),
                Source(
                    workspace_id=workspace.id,
                    name="superseded",
                    source_type=SourceType.TEXT,
                    collection_id="superseded",
                    status=SourceStatus.READY,
                    superseded_at=datetime.now(UTC),
                    ingestion_config={},
                ),
                Source(
                    workspace_id=workspace.id,
                    name="deleted",
                    source_type=SourceType.TEXT,
                    collection_id="deleted",
                    status=SourceStatus.READY,
                    deleted_at=datetime.now(UTC),
                    ingestion_config={},
                ),
                Source(
                    workspace_id=workspace.id,
                    name="missing collection",
                    source_type=SourceType.TEXT,
                    status=SourceStatus.READY,
                    ingestion_config={},
                ),
                Source(
                    workspace_id=other.id,
                    name="other",
                    source_type=SourceType.TEXT,
                    collection_id="other",
                    status=SourceStatus.READY,
                    ingestion_config={},
                ),
            ]
        )
        session.commit()

        sources = SourceRepository(session).ready_collection_ids(workspace.id)

    assert list(sources) == ["active"]
    assert sources["active"].id == active.id


def test_historical_lookup_explicitly_returns_hidden_source_in_its_workspace(
    source_context,
) -> None:
    session_factory, _, workspace, _ = source_context
    with session_factory() as session:
        other = Workspace(title="Other", dataset_id="dataset-other")
        superseded = Source(
            workspace_id=workspace.id,
            name="version 1",
            source_type=SourceType.TEXT,
            collection_id="collection-v1",
            status=SourceStatus.READY,
            ingestion_config={},
        )
        deleted = Source(
            workspace_id=workspace.id,
            name="deleted",
            source_type=SourceType.TEXT,
            collection_id="collection-deleted",
            status=SourceStatus.READY,
            ingestion_config={},
        )
        session.add_all([other, superseded, deleted])
        session.commit()
        repository = SourceRepository(session)
        repository.mark_superseded(superseded.id)
        repository.mark_deleted(deleted.id)

        exists, visible = repository.list_for_workspace(workspace.id)

        assert exists is True
        assert visible == []
        assert repository.get_for_workspace(workspace.id, superseded.id) is None
        historical = repository.get_historical_for_workspace(
            workspace.id, superseded.id
        )
        deleted_history = repository.get_historical_for_workspace(
            workspace.id, deleted.id
        )
        assert historical is not None
        assert historical.summary.id == superseded.id
        assert historical.summary.superseded_at is not None
        assert historical.collection_id == "collection-v1"
        assert deleted_history is not None
        assert deleted_history.summary.id == deleted.id
        assert deleted_history.summary.deleted_at is not None
        assert repository.get_historical_for_workspace(other.id, superseded.id) is None


@pytest.mark.asyncio
async def test_file_ingestion_forwards_original_bytes_and_canonical_config(source_context) -> None:
    _, fake, workspace, service = source_context
    content = b"%PDF-binary-original"
    settings = ChunkSettings.model_validate(
        {
            "trainingType": "qa",
            "indexPrefixTitle": False,
            "customPdfParse": True,
            "chunkSettingMode": "custom",
            "chunkSplitMode": "size",
            "chunkSize": 600,
            "indexSize": 120,
            "chunkSplitter": "",
            "qaPrompt": "Generate grounded pairs.",
        }
    )

    result = await service.ingest_file(
        workspace_id=workspace.id,
        filename="course.pdf",
        content=content,
        settings=settings,
    )

    dataset_id, filename, sent_content, config = fake.create_file_collection_calls[0]
    marker = source_marker(result.source.id)
    assert (dataset_id, filename, sent_content) == (
        "dataset-statistics",
        f"{marker}--course.pdf",
        content,
    )
    assert config == {**settings.model_dump(by_alias=True), "tags": [marker]}
    assert "datasetId" not in config
    assert "collectionId" not in config
    assert result.source.ingestion_config == settings.model_dump(by_alias=True)
    assert result.source.name == "course.pdf"


@pytest.mark.asyncio
async def test_pptx_ingestion_normalizes_locator_records_as_text(source_context) -> None:
    _, fake, workspace, service = source_context
    content = (Path(__file__).parents[1] / "fixtures" / "slides.pptx").read_bytes()

    result = await service.ingest_file(
        workspace_id=workspace.id,
        filename="course.pptx",
        content=content,
        settings=ChunkSettings(),
    )

    assert fake.create_file_collection_calls == []
    dataset_id, remote_name, normalized, _ = fake.create_text_collection_calls[0]
    marker = source_marker(result.source.id)
    assert dataset_id == "dataset-statistics"
    assert remote_name == f"{marker}--course.pptx"
    assert normalized == (
        '[[GT_LOCATOR {"kind":"pptx","slide":1,"title":"Chunking"}]]\n'
        "Chunking\nSplit source material into bounded records.\n\n"
        '[[GT_LOCATOR {"kind":"pptx","slide":2,"title":"Grounding"}]]\n'
        "Grounding\nCitations connect answers to evidence."
    )
    assert result.source.name == "course.pptx"
    assert result.source.source_type is SourceType.FILE


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("content_kind", "expected_text"),
    [
        ("table", "Mean\nAverage"),
        ("nested_group", "Grouped evidence"),
    ],
)
async def test_pptx_ingestion_normalizes_table_and_nested_group_records(
    source_context,
    content_kind: str,
    expected_text: str,
) -> None:
    _, fake, workspace, service = source_context
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    if content_kind == "table":
        table = slide.shapes.add_table(
            1, 2, Inches(1), Inches(1), Inches(6), Inches(1)
        ).table
        table.cell(0, 0).text = "Mean"
        table.cell(0, 1).text = "Average"
    else:
        outer = slide.shapes.add_group_shape()
        inner = outer.shapes.add_group_shape()
        inner.shapes.add_textbox(
            Inches(1), Inches(1), Inches(4), Inches(1)
        ).text = "Grouped evidence"
    output = BytesIO()
    presentation.save(output)

    await service.ingest_file(
        workspace_id=workspace.id,
        filename="nested.pptx",
        content=output.getvalue(),
        settings=ChunkSettings(),
    )

    assert fake.create_text_collection_calls[-1][2] == (
        '[[GT_LOCATOR {"kind":"pptx","slide":1}]]\n' + expected_text
    )


@pytest.mark.asyncio
async def test_xlsx_file_reprocess_uses_normalized_text(source_context) -> None:
    _, fake, workspace, service = source_context
    original = await service.ingest_text(
        workspace_id=workspace.id,
        name="scores",
        text="Original scores",
        settings=ChunkSettings(),
    )
    await service.accept(workspace.id, original.source.id)
    content = (Path(__file__).parents[1] / "fixtures" / "workbook.xlsx").read_bytes()

    result = await service.reprocess_file(
        workspace_id=workspace.id,
        source_id=original.source.id,
        filename="scores.xlsx",
        content=content,
        settings=ChunkSettings(),
    )

    assert fake.create_text_collection_calls[-1][2] == (
        '[[GT_LOCATOR {"cell_range":"A1:B3","kind":"xlsx","sheet":"Week 1"}]]\n'
        "Topic | Score\nMean | 90\nMedian | 85"
    )
    assert result.source.source_type is SourceType.FILE
    assert result.source.replaces_source_id == original.source.id


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("case", "expected_error"),
    [
        ("missing_workspace", SourceWorkspaceNotFoundError),
        ("missing_source", SourceNotFoundError),
        ("invalid_transition", SourceLifecycleConflictError),
        ("pending_review", SourceLifecycleConflictError),
        ("busy", WorkspaceIngestionBusyError),
    ],
)
async def test_file_normalization_runs_only_after_lifecycle_guards(
    source_context,
    monkeypatch: pytest.MonkeyPatch,
    case: str,
    expected_error: type[Exception],
) -> None:
    session_factory, _, workspace, service = source_context
    normalize_calls: list[str] = []

    def normalize(*_args, **_kwargs) -> NormalizedMaterial:
        normalize_calls.append("called")
        return NormalizedMaterial(remote_kind="file", content=b"normalized")

    monkeypatch.setattr("grounded_tutor.services.sources._safe_normalize_material", normalize)
    workspace_id = workspace.id
    source_id = UUID(int=0)
    if case in {"invalid_transition", "pending_review"}:
        source_id = uuid4()
        with session_factory() as session:
            original = Source(
                id=source_id,
                workspace_id=workspace.id,
                name="slides.pptx",
                source_type=SourceType.FILE,
                collection_id="collection-original",
                status=(
                    SourceStatus.INDEXING
                    if case == "invalid_transition"
                    else SourceStatus.READY
                ),
                lineage_id=source_id,
                ingestion_config={},
            )
            session.add(original)
            if case == "pending_review":
                session.add(
                    Source(
                        workspace_id=workspace.id,
                        name="slides v2.pptx",
                        source_type=SourceType.FILE,
                        collection_id="collection-review",
                        status=SourceStatus.REVIEW,
                        lineage_id=source_id,
                        replaces_source_id=source_id,
                        version=2,
                        ingestion_config={},
                    )
                )
            session.commit()
    elif case == "missing_workspace":
        workspace_id = UUID(int=0)

    async def run() -> None:
        if case in {"missing_workspace", "busy"}:
            await service.ingest_file(
                workspace_id=workspace_id,
                filename="slides.pptx",
                content=b"parser input",
                settings=ChunkSettings(),
            )
        else:
            await service.reprocess_file(
                workspace_id=workspace_id,
                source_id=source_id,
                filename="slides.pptx",
                content=b"parser input",
                settings=ChunkSettings(),
            )

    with pytest.raises(expected_error):
        if case == "busy":
            async with service._locks.acquire(workspace.id):
                await run()
        else:
            await run()

    assert normalize_calls == []


@pytest.mark.asyncio
async def test_file_normalization_runs_off_the_event_loop_thread(
    source_context,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _, _, workspace, service = source_context
    event_loop_thread = get_ident()
    normalization_threads: list[int] = []

    def normalize(*_args, **_kwargs) -> NormalizedMaterial:
        normalization_threads.append(get_ident())
        return NormalizedMaterial(remote_kind="file", content=b"normalized")

    monkeypatch.setattr("grounded_tutor.services.sources._safe_normalize_material", normalize)

    await service.ingest_file(
        workspace_id=workspace.id,
        filename="slides.pptx",
        content=b"parser input",
        settings=ChunkSettings(),
    )

    assert normalization_threads
    assert normalization_threads[0] != event_loop_thread


@pytest.mark.asyncio
async def test_cancelled_file_normalization_holds_workspace_lock_until_worker_finishes(
    source_context,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    session_factory, fake, workspace, service = source_context
    started = Event()
    release = Event()
    second_started = Event()
    normalization_calls = 0

    def normalize(*_args, **_kwargs) -> NormalizedMaterial:
        nonlocal normalization_calls
        normalization_calls += 1
        if normalization_calls == 1:
            started.set()
            assert release.wait(timeout=5)
        else:
            second_started.set()
        return NormalizedMaterial(remote_kind="file", content=b"normalized")

    monkeypatch.setattr("grounded_tutor.services.sources._safe_normalize_material", normalize)
    first = asyncio.create_task(
        service.ingest_file(
            workspace_id=workspace.id,
            filename="first.pptx",
            content=b"parser input",
            settings=ChunkSettings(),
        )
    )
    cancelled = False
    try:
        assert await asyncio.wait_for(asyncio.to_thread(started.wait), timeout=1)
        first.cancel()
        await asyncio.sleep(0)
        first.cancel()
        await asyncio.sleep(0)

        assert service._locks.active_count == 1
        with pytest.raises(WorkspaceIngestionBusyError):
            await service.ingest_file(
                workspace_id=workspace.id,
                filename="second.pptx",
                content=b"parser input",
                settings=ChunkSettings(),
            )
        assert not second_started.is_set()
    finally:
        release.set()
        try:
            await first
        except asyncio.CancelledError:
            cancelled = True

    assert cancelled
    assert service._locks.active_count == 0
    assert fake.call_history == []
    with session_factory() as session:
        assert session.scalar(select(Source)) is None


@pytest.mark.asyncio
async def test_image_ingestion_is_gated_and_forwards_verified_original_bytes(
    source_context,
) -> None:
    session_factory, fake, workspace, disabled_service = source_context
    content = (Path(__file__).parents[1] / "fixtures" / "diagram.png").read_bytes()

    with pytest.raises(PreviewError) as caught:
        await disabled_service.ingest_file(
            workspace_id=workspace.id,
            filename="diagram.png",
            content=content,
            settings=ChunkSettings(),
        )

    assert caught.value.code == "unsupported_file_type"
    session = session_factory()
    try:
        enabled_service = SourceService(
            SourceRepository(session),
            fake,
            WorkspaceLockRegistry(),
            supports_image_files=True,
        )
        result = await enabled_service.ingest_file(
            workspace_id=workspace.id,
            filename="diagram.png",
            content=content,
            settings=ChunkSettings(),
        )
    finally:
        session.close()

    assert fake.create_file_collection_calls[-1][2] == content
    assert result.source.name == "diagram.png"


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "marker_like_text",
    [
        '[[GT_LOCATOR {"kind":"pptx","slide":999}]]',
        'x[[GT_LOCATOR{"kind":"pptx","slide":5}]]',
        '[[GT_LOCATOR\t{"kind":"pptx","slide":6}]]',
    ],
)
async def test_normalization_escapes_marker_like_office_content(
    source_context, marker_like_text: str
) -> None:
    _, fake, workspace, service = source_context
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Safety"
    slide.placeholders[1].text = f"{marker_like_text}\nUntrusted content"
    output = BytesIO()
    presentation.save(output)

    await service.ingest_file(
        workspace_id=workspace.id,
        filename="safety.pptx",
        content=output.getvalue(),
        settings=ChunkSettings(),
    )

    normalized = fake.create_text_collection_calls[0][2]
    assert normalized.count("[[GT_LOCATOR") == 1
    assert "[[GT_USER_TEXT" in normalized


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("pptx_title", "expected_title"),
    [
        ('\\"' * 300, ('\\"' * 300)[:120]),
        ("A" * 119 + " " + "B", "A" * 119),
        (
            "Before ]] and [[GT_LOCATOR after",
            "Before ]] and [[GT_LOCATOR after",
        ),
        ("[" * 120, "[" * 120),
    ],
)
async def test_bounded_pptx_title_keeps_preview_and_grounding_locator_coherent(
    source_context, pptx_title: str, expected_title: str
) -> None:
    _, fake, workspace, service = source_context
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = pptx_title
    slide.placeholders[1].text = "Evidence body"
    output = BytesIO()
    presentation.save(output)
    content = output.getvalue()

    preview = preview_file(
        "long-title.pptx",
        content,
        ChunkSettings(),
        max_upload_bytes=20_000_000,
    )
    expected_locator = PptxLocator(slide=1, title=expected_title)
    assert preview.items[0].locator == expected_locator

    ingested = await service.ingest_file(
        workspace_id=workspace.id,
        filename="long-title.pptx",
        content=content,
        settings=ChunkSettings(),
    )
    normalized = fake.create_text_collection_calls[-1][2]
    assert len(normalized.splitlines()[0]) <= MAX_LOCATOR_MARKER_CHARS
    ready_source = await service.accept(workspace.id, ingested.source.id)
    ready = ReadyChunk(
        chunk=RetrievedChunk(
            "chunk-1", "collection-1", "provider", normalized, "", 0.9
        ),
        source=ready_source,
        retrieval_position=1,
    )

    grounded = ground_generated_answer(
        GeneratedAnswer(
            blocks=(
                GeneratedBlock(
                    id="block-1", kind="answer", text="Answer", chunk_ids=("chunk-1",)
                ),
            )
        ),
        {"chunk-1": ready},
        allowed_kinds={"answer"},
    )

    assert grounded.citations[0].locator == expected_locator
    assert "Evidence body" in grounded.citations[0].excerpt


@pytest.mark.asyncio
async def test_office_parser_failures_cross_the_safe_ingestion_boundary(
    source_context,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    session_factory, fake, workspace, service = source_context
    monkeypatch.setattr(
        "grounded_tutor.services.previews.Presentation",
        lambda stream: (_ for _ in ()).throw(RuntimeError("private parser detail")),
    )

    with pytest.raises(PreviewError) as caught:
        await service.ingest_file(
            workspace_id=workspace.id,
            filename="slides.pptx",
            content=(Path(__file__).parents[1] / "fixtures" / "slides.pptx").read_bytes(),
            settings=ChunkSettings(),
        )

    assert caught.value.code == "unreadable_file"
    assert "private parser detail" not in str(caught.value)
    assert fake.call_history == []
    with session_factory() as session:
        assert session.scalar(select(Source)) is None


@pytest.mark.asyncio
async def test_failure_before_collection_persists_retryable_failed_source(source_context) -> None:
    session_factory, fake, workspace, service = source_context
    fake.failures["create_text_collection"] = ExternalServiceError(
        service="fastgpt", category="timeout", safe_message="private remote detail"
    )
    settings = ChunkSettings(trainingType="qa", qaPrompt="retry me")

    with pytest.raises(ExternalSourceServiceError):
        await service.ingest_text(
            workspace_id=workspace.id,
            name="notes",
            text="Mean.",
            settings=settings,
        )

    with session_factory() as session:
        source = session.scalar(select(Source))
    assert source is not None
    assert source.status is SourceStatus.FAILED
    assert source.collection_id is None
    assert source.ingestion_config == settings.model_dump(by_alias=True)
    assert source.error_message == "Source ingestion failed."
    assert "private" not in source.error_message


@pytest.mark.asyncio
async def test_failure_after_collection_keeps_exact_collection_forbidden(source_context) -> None:
    session_factory, fake, workspace, service = source_context
    fake.failures["list_collection_data"] = RuntimeError("secret content and api-key")

    with pytest.raises(ExternalSourceServiceError):
        await service.ingest_text(
            workspace_id=workspace.id,
            name="notes",
            text="Mean.",
            settings=ChunkSettings(),
        )

    with session_factory() as session:
        source = session.scalar(select(Source))
    assert source is not None
    assert source.status is SourceStatus.FAILED
    assert source.collection_id == "collection-1"
    assert fake.collections["collection-1"].forbidden is True
    assert fake.set_collection_forbidden_calls == [
        ("collection-1", True),
        ("collection-1", True),
    ]
    assert "secret" not in (source.error_message or "")


@pytest.mark.asyncio
async def test_same_workspace_ingestion_is_rejected_without_waiting(
    source_engine,
) -> None:
    with Session(source_engine, expire_on_commit=False) as setup:
        workspace = Workspace(title="Statistics", dataset_id="dataset-statistics")
        setup.add(workspace)
        setup.commit()
        setup.refresh(workspace)

    entered = asyncio.Event()
    release = asyncio.Event()

    class BlockingFake(FakeFastGPT):
        async def create_text_collection(self, dataset_id, name, text, config):
            entered.set()
            await release.wait()
            return await super().create_text_collection(dataset_id, name, text, config)

    fake = BlockingFake()
    fake.datasets[workspace.dataset_id] = DatasetRef(workspace.dataset_id)
    locks = WorkspaceLockRegistry()
    first_session = Session(source_engine, expire_on_commit=False)
    second_session = Session(source_engine, expire_on_commit=False)
    first = SourceService(SourceRepository(first_session), fake, locks)
    second = SourceService(SourceRepository(second_session), fake, locks)
    try:
        active = asyncio.create_task(
            first.ingest_text(
                workspace_id=workspace.id,
                name="first",
                text="Mean.",
                settings=ChunkSettings(),
            )
        )
        await asyncio.wait_for(entered.wait(), timeout=1)
        with pytest.raises(WorkspaceIngestionBusyError):
            await asyncio.wait_for(
                second.ingest_text(
                    workspace_id=workspace.id,
                    name="second",
                    text="Median.",
                    settings=ChunkSettings(),
                ),
                timeout=0.1,
            )
        release.set()
        await active
    finally:
        first_session.close()
        second_session.close()

    assert locks.active_count == 0
    assert len(fake.create_text_collection_calls) == 1


@pytest.mark.asyncio
async def test_different_workspaces_may_ingest_concurrently(source_engine) -> None:
    with Session(source_engine, expire_on_commit=False) as setup:
        workspaces = [
            Workspace(title="Statistics", dataset_id="dataset-statistics"),
            Workspace(title="Calculus", dataset_id="dataset-calculus"),
        ]
        setup.add_all(workspaces)
        setup.commit()
        for workspace in workspaces:
            setup.refresh(workspace)

    both_entered = asyncio.Event()
    release = asyncio.Event()
    entered: set[str] = set()

    class BlockingFake(FakeFastGPT):
        async def create_text_collection(self, dataset_id, name, text, config):
            entered.add(dataset_id)
            if len(entered) == 2:
                both_entered.set()
            await release.wait()
            return await super().create_text_collection(dataset_id, name, text, config)

    fake = BlockingFake()
    for workspace in workspaces:
        fake.datasets[workspace.dataset_id] = DatasetRef(workspace.dataset_id)
    locks = WorkspaceLockRegistry()
    sessions = [Session(source_engine, expire_on_commit=False) for _ in workspaces]
    services = [SourceService(SourceRepository(session), fake, locks) for session in sessions]
    try:
        tasks = [
            asyncio.create_task(
                service.ingest_text(
                    workspace_id=workspace.id,
                    name=workspace.title,
                    text="Content.",
                    settings=ChunkSettings(),
                )
            )
            for service, workspace in zip(services, workspaces, strict=True)
        ]
        await asyncio.wait_for(both_entered.wait(), timeout=1)
        release.set()
        await asyncio.gather(*tasks)
    finally:
        for session in sessions:
            session.close()

    assert locks.active_count == 0
    assert entered == {"dataset-statistics", "dataset-calculus"}


@pytest.mark.asyncio
async def test_cancellation_releases_workspace_lock(source_engine) -> None:
    with Session(source_engine, expire_on_commit=False) as setup:
        workspace = Workspace(title="Statistics", dataset_id="dataset-statistics")
        setup.add(workspace)
        setup.commit()
        setup.refresh(workspace)

    entered = asyncio.Event()

    class NeverReturnsFake(FakeFastGPT):
        async def create_text_collection(self, dataset_id, name, text, config):
            entered.set()
            await asyncio.Event().wait()
            raise AssertionError("unreachable")

    fake = NeverReturnsFake()
    fake.datasets[workspace.dataset_id] = DatasetRef(workspace.dataset_id)
    locks = WorkspaceLockRegistry()
    session = Session(source_engine, expire_on_commit=False)
    service = SourceService(SourceRepository(session), fake, locks)
    try:
        task = asyncio.create_task(
            service.ingest_text(
                workspace_id=workspace.id,
                name="notes",
                text="Mean.",
                settings=ChunkSettings(),
            )
        )
        await asyncio.wait_for(entered.wait(), timeout=1)
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task
    finally:
        session.close()

    assert locks.active_count == 0


@pytest.mark.asyncio
async def test_ambiguous_indexing_commit_is_conservatively_marked_failed(
    source_engine,
) -> None:
    with Session(source_engine, expire_on_commit=False) as setup:
        workspace = Workspace(title="Statistics", dataset_id="dataset-statistics")
        setup.add(workspace)
        setup.commit()
        setup.refresh(workspace)

    class FirstAfterCommitFails(Session):
        def __init__(self, *args, **kwargs) -> None:
            super().__init__(*args, **kwargs)
            self.failed = False
            sqlalchemy_event.listen(self, "after_commit", self._fail_once)

        def _fail_once(self, session) -> None:
            if not self.failed:
                self.failed = True
                raise SQLAlchemyError("ambiguous private database detail")

    fake = FakeFastGPT()
    fake.datasets[workspace.dataset_id] = DatasetRef(workspace.dataset_id)
    session = FirstAfterCommitFails(bind=source_engine, expire_on_commit=False)
    service = SourceService(SourceRepository(session), fake, WorkspaceLockRegistry())
    try:
        with pytest.raises(Exception) as caught:
            await service.ingest_text(
                workspace_id=workspace.id,
                name="notes",
                text="Mean.",
                settings=ChunkSettings(),
            )
    finally:
        session.close()

    assert caught.type.__name__ == "SourcePersistenceError"
    assert fake.call_history == []
    with Session(source_engine) as check:
        source = check.scalar(select(Source))
    assert source is not None
    assert source.status is SourceStatus.FAILED
    assert source.error_message == "Source ingestion failed."


@pytest.mark.asyncio
async def test_collection_id_flush_failure_leaves_remote_disabled_and_local_failed(
    source_engine,
) -> None:
    with Session(source_engine, expire_on_commit=False) as setup:
        workspace = Workspace(title="Statistics", dataset_id="dataset-statistics")
        setup.add(workspace)
        setup.commit()
        setup.refresh(workspace)

    class CollectionFlushFailsOnce(Session):
        def __init__(self, *args, **kwargs) -> None:
            super().__init__(*args, **kwargs)
            self.failed = False
            self.source_flush_calls = 0

        def flush(self, objects=None) -> None:
            if any(isinstance(instance, Source) for instance in self.new | self.dirty):
                self.source_flush_calls += 1
            if self.source_flush_calls == 2 and not self.failed:
                self.failed = True
                raise SQLAlchemyError("private collection write detail")
            super().flush(objects)

    fake = FakeFastGPT()
    fake.datasets[workspace.dataset_id] = DatasetRef(workspace.dataset_id)
    session = CollectionFlushFailsOnce(bind=source_engine, expire_on_commit=False)
    service = SourceService(SourceRepository(session), fake, WorkspaceLockRegistry())
    try:
        with pytest.raises(Exception) as caught:
            await service.ingest_text(
                workspace_id=workspace.id,
                name="notes",
                text="Mean.",
                settings=ChunkSettings(),
            )
    finally:
        session.close()

    assert caught.type.__name__ == "SourcePersistenceError"
    assert fake.collections["collection-1"].forbidden is True
    assert fake.set_collection_forbidden_calls == [
        ("collection-1", True),
        ("collection-1", True),
    ]
    with Session(source_engine) as check:
        source = check.scalar(select(Source))
    assert source is not None
    assert source.status is SourceStatus.FAILED
    assert source.collection_id == "collection-1"


@pytest.mark.asyncio
async def test_forbid_and_compensation_failures_return_safe_failure_state(
    source_context,
) -> None:
    session_factory, fake, workspace, service = source_context
    fake.set_collection_forbidden = AsyncMock(
        side_effect=RuntimeError("collection-1 private api-key")
    )

    with pytest.raises(ExternalSourceServiceError) as caught:
        await service.ingest_text(
            workspace_id=workspace.id,
            name="notes",
            text="Mean.",
            settings=ChunkSettings(),
        )

    assert "private" not in str(caught.value)
    with session_factory() as session:
        source = session.scalar(select(Source))
    assert source is not None
    assert source.status is SourceStatus.FAILED
    assert source.collection_id == "collection-1"
    assert fake.set_collection_forbidden.await_count == 3


@pytest.mark.asyncio
async def test_initial_source_add_failure_never_calls_fastgpt(source_engine) -> None:
    with Session(source_engine, expire_on_commit=False) as setup:
        workspace = Workspace(title="Statistics", dataset_id="dataset-statistics")
        setup.add(workspace)
        setup.commit()
        setup.refresh(workspace)

    class SourceAddFails(Session):
        def add(self, instance, _warn: bool = True) -> None:
            if isinstance(instance, Source):
                raise SQLAlchemyError("private add detail")
            super().add(instance, _warn=_warn)

    fake = FakeFastGPT()
    fake.datasets[workspace.dataset_id] = DatasetRef(workspace.dataset_id)
    session = SourceAddFails(bind=source_engine, expire_on_commit=False)
    service = SourceService(SourceRepository(session), fake, WorkspaceLockRegistry())
    try:
        with pytest.raises(Exception) as caught:
            await service.ingest_text(
                workspace_id=workspace.id,
                name="notes",
                text="Mean.",
                settings=ChunkSettings(),
            )
    finally:
        session.close()

    assert caught.type.__name__ == "SourcePersistenceError"
    assert fake.call_history == []
    with Session(source_engine) as check:
        assert check.scalar(select(Source)) is None


@pytest.mark.asyncio
async def test_review_flush_failure_is_recovered_to_failed(source_engine) -> None:
    with Session(source_engine, expire_on_commit=False) as setup:
        workspace = Workspace(title="Statistics", dataset_id="dataset-statistics")
        setup.add(workspace)
        setup.commit()
        setup.refresh(workspace)

    class ReviewFlushFailsOnce(Session):
        def __init__(self, *args, **kwargs) -> None:
            super().__init__(*args, **kwargs)
            self.source_flush_calls = 0

        def flush(self, objects=None) -> None:
            if any(isinstance(instance, Source) for instance in self.new | self.dirty):
                self.source_flush_calls += 1
            if self.source_flush_calls == 3:
                self.source_flush_calls += 1
                raise SQLAlchemyError("private review detail")
            super().flush(objects)

    fake = FakeFastGPT()
    fake.datasets[workspace.dataset_id] = DatasetRef(workspace.dataset_id)
    session = ReviewFlushFailsOnce(bind=source_engine, expire_on_commit=False)
    service = SourceService(SourceRepository(session), fake, WorkspaceLockRegistry())
    try:
        with pytest.raises(Exception) as caught:
            await service.ingest_text(
                workspace_id=workspace.id,
                name="notes",
                text="Mean.",
                settings=ChunkSettings(),
            )
    finally:
        session.close()

    assert caught.type.__name__ == "SourcePersistenceError"
    assert fake.collections["collection-1"].forbidden is True
    with Session(source_engine) as check:
        source = check.scalar(select(Source))
    assert source is not None
    assert source.status is SourceStatus.FAILED
    assert source.collection_id == "collection-1"


@pytest.mark.asyncio
async def test_failed_transition_db_error_does_not_expose_external_error(source_engine) -> None:
    with Session(source_engine, expire_on_commit=False) as setup:
        workspace = Workspace(title="Statistics", dataset_id="dataset-statistics")
        setup.add(workspace)
        setup.commit()
        setup.refresh(workspace)

    class FailedTransitionFlushFails(Session):
        def flush(self, objects=None) -> None:
            if any(
                isinstance(instance, Source) and instance.status is SourceStatus.FAILED
                for instance in self.dirty
            ):
                raise SQLAlchemyError("private failed transition detail")
            super().flush(objects)

    fake = FakeFastGPT()
    fake.datasets[workspace.dataset_id] = DatasetRef(workspace.dataset_id)
    fake.failures["create_text_collection"] = RuntimeError("external private material")
    session = FailedTransitionFlushFails(bind=source_engine, expire_on_commit=False)
    service = SourceService(SourceRepository(session), fake, WorkspaceLockRegistry())
    try:
        with pytest.raises(Exception) as caught:
            await service.ingest_text(
                workspace_id=workspace.id,
                name="notes",
                text="Mean.",
                settings=ChunkSettings(),
            )
    finally:
        session.close()

    assert caught.type.__name__ == "SourcePersistenceError"
    assert "private" not in str(caught.value)
    assert fake.collections == {}


@pytest.mark.asyncio
async def test_remote_create_then_timeout_reconciles_enabled_orphan(source_engine) -> None:
    with Session(source_engine, expire_on_commit=False) as setup:
        workspace = Workspace(title="Statistics", dataset_id="dataset-statistics")
        setup.add(workspace)
        setup.commit()
        setup.refresh(workspace)

    class CommitThenTimeout(FakeFastGPT):
        async def create_text_collection(self, dataset_id, name, text, config):
            await super().create_text_collection(dataset_id, name, text, config)
            raise ExternalServiceError(
                service="fastgpt",
                category="timeout",
                safe_message="FastGPT request failed.",
            )

    fake = CommitThenTimeout()
    fake.datasets[workspace.dataset_id] = DatasetRef(workspace.dataset_id)
    session = Session(source_engine, expire_on_commit=False)
    service = SourceService(SourceRepository(session), fake, WorkspaceLockRegistry())
    try:
        with pytest.raises(ExternalSourceServiceError):
            await service.ingest_text(
                workspace_id=workspace.id,
                name="notes",
                text="Mean.",
                settings=ChunkSettings(),
            )
    finally:
        session.close()

    assert len(fake.collections) == 1
    remote = next(iter(fake.collections.values()))
    assert remote.forbidden is True
    with Session(source_engine) as check:
        source = check.scalar(select(Source))
    assert source is not None
    assert source.status is SourceStatus.FAILED
    assert source.collection_id is not None
    assert source.name == "notes"


@pytest.mark.asyncio
async def test_remote_create_then_cancellation_reconciles_enabled_orphan(source_engine) -> None:
    with Session(source_engine, expire_on_commit=False) as setup:
        workspace = Workspace(title="Statistics", dataset_id="dataset-statistics")
        setup.add(workspace)
        setup.commit()
        setup.refresh(workspace)

    committed = asyncio.Event()

    class CommitThenBlock(FakeFastGPT):
        async def create_text_collection(self, dataset_id, name, text, config):
            await super().create_text_collection(dataset_id, name, text, config)
            committed.set()
            await asyncio.Event().wait()
            raise AssertionError("unreachable")

    fake = CommitThenBlock()
    fake.datasets[workspace.dataset_id] = DatasetRef(workspace.dataset_id)
    session = Session(source_engine, expire_on_commit=False)
    service = SourceService(SourceRepository(session), fake, WorkspaceLockRegistry())
    try:
        task = asyncio.create_task(
            service.ingest_text(
                workspace_id=workspace.id,
                name="notes",
                text="Mean.",
                settings=ChunkSettings(),
            )
        )
        await asyncio.wait_for(committed.wait(), timeout=1)
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task
    finally:
        session.close()

    assert len(fake.collections) == 1
    assert next(iter(fake.collections.values())).forbidden is True
    with Session(source_engine) as check:
        source = check.scalar(select(Source))
    assert source is not None
    assert source.status is SourceStatus.FAILED
    assert source.collection_id is not None


@pytest.mark.asyncio
async def test_name_marker_recovers_when_collection_tags_are_omitted(source_engine) -> None:
    with Session(source_engine, expire_on_commit=False) as setup:
        workspace = Workspace(title="Statistics", dataset_id="dataset-statistics")
        setup.add(workspace)
        setup.commit()
        setup.refresh(workspace)

    class CommitThenMalformed(FakeFastGPT):
        async def create_text_collection(self, dataset_id, name, text, config):
            await super().create_text_collection(dataset_id, name, text, config)
            raise ExternalServiceError(
                service="fastgpt",
                category="malformed_response",
                safe_message="FastGPT returned a malformed response.",
            )

    fake = CommitThenMalformed()
    fake.supports_tags = False
    fake.datasets[workspace.dataset_id] = DatasetRef(workspace.dataset_id)
    session = Session(source_engine, expire_on_commit=False)
    service = SourceService(SourceRepository(session), fake, WorkspaceLockRegistry())
    try:
        with pytest.raises(ExternalSourceServiceError):
            await service.ingest_text(
                workspace_id=workspace.id,
                name="notes",
                text="Mean.",
                settings=ChunkSettings(),
            )
    finally:
        session.close()

    remote = next(iter(fake.collections.values()))
    assert remote.tags == ()
    assert remote.forbidden is True
    with Session(source_engine) as check:
        source = check.scalar(select(Source))
    assert source is not None
    assert source.collection_id == "collection-1"


@pytest.mark.asyncio
async def test_reconciliation_ignores_wrong_dataset_and_near_marker_matches(
    source_engine,
) -> None:
    with Session(source_engine, expire_on_commit=False) as setup:
        workspace = Workspace(title="Statistics", dataset_id="dataset-statistics")
        setup.add(workspace)
        setup.commit()
        setup.refresh(workspace)

    class AddsDecoysThenTimesOut(FakeFastGPT):
        async def create_text_collection(self, dataset_id, name, text, config):
            ref = await super().create_text_collection(dataset_id, name, text, config)
            marker = name.split("--", maxsplit=1)[0]
            self.collections["near-match"] = FakeCollection(
                dataset_id=dataset_id,
                name=f"{marker}x--near",
                content="near",
                config={},
                chunks=[],
                tags=(f"{marker}x",),
            )
            self.collections["wrong-dataset"] = FakeCollection(
                dataset_id="dataset-other",
                name=name,
                content="wrong",
                config={},
                chunks=[],
                tags=(marker,),
            )
            assert ref.collection_id == "collection-1"
            raise ExternalServiceError(
                service="fastgpt",
                category="timeout",
                safe_message="FastGPT request failed.",
            )

    fake = AddsDecoysThenTimesOut()
    fake.datasets[workspace.dataset_id] = DatasetRef(workspace.dataset_id)
    fake.datasets["dataset-other"] = DatasetRef("dataset-other")
    session = Session(source_engine, expire_on_commit=False)
    service = SourceService(SourceRepository(session), fake, WorkspaceLockRegistry())
    try:
        with pytest.raises(ExternalSourceServiceError):
            await service.ingest_text(
                workspace_id=workspace.id,
                name="notes",
                text="Mean.",
                settings=ChunkSettings(),
            )
    finally:
        session.close()

    assert fake.collections["collection-1"].forbidden is True
    assert fake.collections["near-match"].forbidden is False
    assert fake.collections["wrong-dataset"].forbidden is False
    assert all(
        call[0] == "dataset-statistics" for call in fake.list_collections_calls
    )


@pytest.mark.asyncio
async def test_multiple_exact_matches_are_all_disabled_and_local_id_is_ambiguous(
    source_engine,
) -> None:
    with Session(source_engine, expire_on_commit=False) as setup:
        workspace = Workspace(title="Statistics", dataset_id="dataset-statistics")
        setup.add(workspace)
        setup.commit()
        setup.refresh(workspace)

    class DuplicatesThenTimesOut(FakeFastGPT):
        async def create_text_collection(self, dataset_id, name, text, config):
            await super().create_text_collection(dataset_id, name, text, config)
            marker = name.split("--", maxsplit=1)[0]
            self.collections["collection-duplicate"] = FakeCollection(
                dataset_id=dataset_id,
                name=f"{marker}--duplicate",
                content="duplicate",
                config=dict(config),
                chunks=[],
                tags=(marker,),
            )
            raise ExternalServiceError(
                service="fastgpt",
                category="timeout",
                safe_message="FastGPT request failed.",
            )

    fake = DuplicatesThenTimesOut()
    fake.datasets[workspace.dataset_id] = DatasetRef(workspace.dataset_id)
    session = Session(source_engine, expire_on_commit=False)
    service = SourceService(SourceRepository(session), fake, WorkspaceLockRegistry())
    try:
        with pytest.raises(ExternalSourceServiceError):
            await service.ingest_text(
                workspace_id=workspace.id,
                name="notes",
                text="Mean.",
                settings=ChunkSettings(),
            )
    finally:
        session.close()

    assert fake.collections["collection-1"].forbidden is True
    assert fake.collections["collection-duplicate"].forbidden is True
    with Session(source_engine) as check:
        source = check.scalar(select(Source))
    assert source is not None
    assert source.status is SourceStatus.FAILED
    assert source.collection_id is None
    assert source.error_message == (
        "Source ingestion failed; multiple remote matches require review."
    )


@pytest.mark.asyncio
async def test_reconciliation_paginates_to_exact_match(source_engine) -> None:
    with Session(source_engine, expire_on_commit=False) as setup:
        workspace = Workspace(title="Statistics", dataset_id="dataset-statistics")
        setup.add(workspace)
        setup.commit()
        setup.refresh(workspace)

    class PaginatesThenTimesOut(FakeFastGPT):
        async def create_text_collection(self, dataset_id, name, text, config):
            marker = name.split("--", maxsplit=1)[0]
            for index in range(30):
                self.collections[f"near-{index}"] = FakeCollection(
                    dataset_id=dataset_id,
                    name=f"{marker}x--near-{index}",
                    content="near",
                    config={},
                    chunks=[],
                )
            await super().create_text_collection(dataset_id, name, text, config)
            raise ExternalServiceError(
                service="fastgpt",
                category="timeout",
                safe_message="FastGPT request failed.",
            )

    fake = PaginatesThenTimesOut()
    fake.datasets[workspace.dataset_id] = DatasetRef(workspace.dataset_id)
    session = Session(source_engine, expire_on_commit=False)
    service = SourceService(SourceRepository(session), fake, WorkspaceLockRegistry())
    try:
        with pytest.raises(ExternalSourceServiceError):
            await service.ingest_text(
                workspace_id=workspace.id,
                name="notes",
                text="Mean.",
                settings=ChunkSettings(),
            )
    finally:
        session.close()

    assert fake.list_collections_calls == [
        ("dataset-statistics", 0, 30, fake.list_collections_calls[0][3]),
        ("dataset-statistics", 30, 30, fake.list_collections_calls[0][3]),
    ]
    assert fake.collections["collection-1"].forbidden is True


@pytest.mark.asyncio
async def test_persistent_list_outage_fails_locally_without_claiming_remote_disable(
    source_engine,
) -> None:
    with Session(source_engine, expire_on_commit=False) as setup:
        workspace = Workspace(title="Statistics", dataset_id="dataset-statistics")
        setup.add(workspace)
        setup.commit()
        setup.refresh(workspace)

    class CommitThenTimeoutWithListOutage(FakeFastGPT):
        async def create_text_collection(self, dataset_id, name, text, config):
            await super().create_text_collection(dataset_id, name, text, config)
            raise ExternalServiceError(
                service="fastgpt",
                category="timeout",
                safe_message="FastGPT request failed.",
            )

        async def list_collections(self, *args, **kwargs):
            raise ExternalServiceError(
                service="fastgpt",
                category="network",
                safe_message="FastGPT request failed.",
            )

    fake = CommitThenTimeoutWithListOutage()
    fake.datasets[workspace.dataset_id] = DatasetRef(workspace.dataset_id)
    session = Session(source_engine, expire_on_commit=False)
    service = SourceService(SourceRepository(session), fake, WorkspaceLockRegistry())
    try:
        with pytest.raises(ExternalSourceServiceError):
            await service.ingest_text(
                workspace_id=workspace.id,
                name="notes",
                text="Mean.",
                settings=ChunkSettings(),
            )
    finally:
        session.close()

    assert fake.collections["collection-1"].forbidden is False
    with Session(source_engine) as check:
        source = check.scalar(select(Source))
    assert source is not None
    assert source.status is SourceStatus.FAILED
    assert source.collection_id is None
    assert source.error_message == (
        "Source ingestion failed; remote reconciliation could not be completed."
    )


@pytest.mark.asyncio
async def test_persistent_forbid_failure_never_transitions_to_review(source_engine) -> None:
    with Session(source_engine, expire_on_commit=False) as setup:
        workspace = Workspace(title="Statistics", dataset_id="dataset-statistics")
        setup.add(workspace)
        setup.commit()
        setup.refresh(workspace)

    class ForbidOutage(FakeFastGPT):
        async def set_collection_forbidden(self, collection_id, forbidden):
            raise ExternalServiceError(
                service="fastgpt",
                category="network",
                safe_message="FastGPT request failed.",
            )

    fake = ForbidOutage()
    fake.datasets[workspace.dataset_id] = DatasetRef(workspace.dataset_id)
    session = Session(source_engine, expire_on_commit=False)
    service = SourceService(SourceRepository(session), fake, WorkspaceLockRegistry())
    try:
        with pytest.raises(ExternalSourceServiceError):
            await service.ingest_text(
                workspace_id=workspace.id,
                name="notes",
                text="Mean.",
                settings=ChunkSettings(),
            )
    finally:
        session.close()

    assert fake.collections["collection-1"].forbidden is False
    with Session(source_engine) as check:
        source = check.scalar(select(Source))
    assert source is not None
    assert source.status is SourceStatus.FAILED
    assert source.collection_id == "collection-1"
    assert source.error_message == (
        "Source ingestion failed; remote disable could not be confirmed."
    )
