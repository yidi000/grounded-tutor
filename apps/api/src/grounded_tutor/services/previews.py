from __future__ import annotations

import csv
import datetime as dt
import re
import struct
import warnings
from collections.abc import Collection, Iterator, Mapping
from io import BytesIO, StringIO
from pathlib import Path
from typing import ClassVar
from uuid import UUID
from zipfile import ZipFile

from bs4 import BeautifulSoup
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PIL import Image, UnidentifiedImageError
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.shapes.base import BaseShape
from pypdf import PdfReader

from grounded_tutor.domain.answers import ImageLocator, PptxLocator, SourceLocator, XlsxLocator
from grounded_tutor.domain.errors import PublicErrorCode
from grounded_tutor.domain.ingestion import (
    ChunkSettings,
    PreviewItem,
    PreviewResponse,
    PreviewWarning,
)
from grounded_tutor.repositories.workspaces import WorkspaceRepository

MAX_PREVIEW_ITEMS = 50
MAX_PREVIEW_EXCERPT_CHARS = 500
# Deterministic P0 guardrails. They bound parser inputs and output work; they do
# not claim to sandbox third-party parsers or provide a wall-clock timeout.
DEFAULT_MAX_EXTRACTED_CHARACTERS = 40_000_000
MAX_ZIP_ARCHIVE_MEMBERS = 1_000
MAX_ZIP_CENTRAL_DIRECTORY_BYTES = 2_000_000
MAX_ZIP_UNCOMPRESSED_BYTES = 80_000_000
MAX_ZIP_COMPRESSION_RATIO = 100.0
MAX_DOCX_BLOCKS = 10_000
MAX_PDF_PAGES = 500
MAX_PPTX_SLIDES = 500
MAX_PPTX_SHAPES = 10_000
MAX_PPTX_LOCATOR_TITLE_CHARS = 120
MAX_XLSX_SHEETS = 100
MAX_XLSX_ROWS = 100_000
MAX_XLSX_CELLS = 1_000_000
MAX_IMAGE_DIMENSION = 16_384
MAX_IMAGE_PIXELS = 40_000_000
IMAGE_EXTENSIONS = frozenset({".png", ".jpg", ".jpeg", ".webp"})
SUPPORTED_EXTENSIONS = frozenset(
    {".pdf", ".docx", ".md", ".txt", ".html", ".csv", ".pptx", ".xlsx"}
)

ZIP_EOCD_SIGNATURE = b"PK\x05\x06"
ZIP64_EOCD_SIGNATURE = b"PK\x06\x06"
ZIP64_LOCATOR_SIGNATURE = b"PK\x06\x07"
ZIP_CENTRAL_DIRECTORY_SIGNATURE = b"PK\x01\x02"
ZIP_LOCAL_FILE_SIGNATURE = b"PK\x03\x04"
ZIP_EOCD_SIZE = 22
ZIP_MAX_COMMENT_BYTES = 65_535
ZIP_CENTRAL_DIRECTORY_HEADER_SIZE = 46
ZIP64_EXTRA_FIELD_ID = 0x0001
ZIP16_SENTINEL = 0xFFFF
ZIP32_SENTINEL = 0xFFFFFFFF


class PreviewError(RuntimeError):
    """Stable, content-free failure exposed by the local preview boundary."""

    _MESSAGES: ClassVar[Mapping[str, str]] = {
        "file_too_large": "The uploaded file exceeds the preview size limit.",
        "text_too_large": "The pasted text exceeds the preview size limit.",
        "source_too_large": "The extracted source text exceeds the preview limit.",
        "source_work_limit_exceeded": "The source exceeds the preview processing limit.",
        "unsafe_archive": "The document archive exceeds safe preview limits.",
        "validation_error": "The source name is invalid.",
        "unsupported_file_type": "This file type is not supported for preview.",
        "empty_source": "The source does not contain readable text.",
        "unreadable_file": "The file could not be read for preview.",
    }

    def __init__(self, code: PublicErrorCode) -> None:
        self.code = code
        super().__init__(self._MESSAGES[code])


class PreviewWorkspaceNotFoundError(LookupError):
    pass


class PreviewService:
    def __init__(
        self,
        workspaces: WorkspaceRepository,
        *,
        max_upload_bytes: int,
        max_preview_text_bytes: int,
        max_extracted_characters: int,
        supports_image_files: bool = False,
    ) -> None:
        _require_positive_limit(max_upload_bytes)
        _require_positive_limit(max_preview_text_bytes)
        _require_positive_limit(max_extracted_characters)
        self._workspaces = workspaces
        self.max_upload_bytes = max_upload_bytes
        self.max_preview_text_bytes = max_preview_text_bytes
        self.max_extracted_characters = max_extracted_characters
        self.supports_image_files = supports_image_files

    def from_text(
        self,
        workspace_id: UUID,
        *,
        source_name: str,
        text: str,
        settings: ChunkSettings,
    ) -> PreviewResponse:
        _validate_text_budget(
            text,
            max_text_bytes=self.max_preview_text_bytes,
            max_characters=self.max_extracted_characters,
        )
        self._require_workspace(workspace_id)
        return preview_text(
            text,
            settings,
            source_name=source_name,
            max_text_bytes=self.max_preview_text_bytes,
            max_extracted_characters=self.max_extracted_characters,
        )

    def from_file(
        self,
        workspace_id: UUID,
        *,
        filename: str,
        content: bytes,
        settings: ChunkSettings,
    ) -> PreviewResponse:
        _validate_file_size(content, self.max_upload_bytes)
        self._require_workspace(workspace_id)
        return preview_file(
            filename,
            content,
            settings,
            max_upload_bytes=self.max_upload_bytes,
            max_extracted_characters=self.max_extracted_characters,
            supports_image_files=self.supports_image_files,
        )

    def _require_workspace(self, workspace_id: UUID) -> None:
        if self._workspaces.get(workspace_id) is None:
            raise PreviewWorkspaceNotFoundError


def preview_text(
    text: str,
    settings: ChunkSettings,
    *,
    source_name: str = "Pasted text",
    max_text_bytes: int | None = None,
    max_extracted_characters: int | None = None,
) -> PreviewResponse:
    _validate_text_budget(
        text,
        max_text_bytes=max_text_bytes,
        max_characters=max_extracted_characters,
    )
    clean_text = text.strip()
    if not clean_text:
        raise PreviewError("empty_source")
    items: list[PreviewItem] = []
    preview_truncated = False
    for position, segment in enumerate(_segments(clean_text, settings), start=1):
        if position > MAX_PREVIEW_ITEMS:
            preview_truncated = True
            break
        items.append(
            PreviewItem(
                position=position,
                text=segment[:MAX_PREVIEW_EXCERPT_CHARS],
                character_count=len(segment),
                truncated=len(segment) > MAX_PREVIEW_EXCERPT_CHARS,
            )
        )
    if not items:
        raise PreviewError("empty_source")
    warnings = []
    if settings.training_type == "qa":
        warnings.append(PreviewWarning(code="qa_generated_after_processing"))
    if preview_truncated:
        warnings.append(PreviewWarning(code="preview_truncated"))
    return PreviewResponse(
        source_name=source_name,
        character_count=len(clean_text),
        items=items,
        warnings=warnings,
    )


def preview_file(
    filename: str,
    content: bytes,
    settings: ChunkSettings,
    *,
    max_upload_bytes: int,
    max_extracted_characters: int = DEFAULT_MAX_EXTRACTED_CHARACTERS,
    supports_image_files: bool = False,
) -> PreviewResponse:
    _require_positive_limit(max_extracted_characters)
    _validate_file_size(content, max_upload_bytes)
    source_name = safe_source_name(filename)
    suffix = Path(source_name).suffix.lower()
    supported_extensions = SUPPORTED_EXTENSIONS | (
        IMAGE_EXTENSIONS if supports_image_files else frozenset()
    )
    if suffix not in supported_extensions or not Path(source_name).stem:
        raise PreviewError("unsupported_file_type")
    try:
        if suffix in IMAGE_EXTENSIONS:
            text = image_metadata(suffix, content)
            return _preview_records(
                source_name,
                [(text, ImageLocator(filename=source_name))],
                settings,
            )
        if suffix in {".pptx", ".xlsx"}:
            return _preview_records(
                source_name,
                extract_office_records(
                    suffix,
                    content,
                    max_extracted_characters=max_extracted_characters,
                ),
                settings,
            )
        text = _extract_text(
            suffix,
            content,
            max_extracted_characters=max_extracted_characters,
        )
    except PreviewError:
        raise
    except Exception:  # noqa: BLE001 -- all third-party parser failures cross a safe boundary
        raise PreviewError("unreadable_file") from None
    return preview_text(
        text,
        settings,
        source_name=source_name,
        max_extracted_characters=max_extracted_characters,
    )


def _extract_text(
    suffix: str,
    content: bytes,
    *,
    max_extracted_characters: int,
) -> str:
    if suffix in {".md", ".txt"}:
        return _bounded_text(content.decode("utf-8-sig"), max_extracted_characters)
    if suffix == ".html":
        soup = BeautifulSoup(content.decode("utf-8-sig"), "html.parser")
        for hidden in soup(["script", "style", "noscript", "template"]):
            hidden.decompose()
        return _bounded_text(
            soup.get_text("\n", strip=True), max_extracted_characters
        )
    if suffix == ".csv":
        rows = csv.reader(
            StringIO(content.decode("utf-8-sig"), newline=""), strict=True
        )
        extracted_rows: list[str] = []
        extracted_characters = 0
        for row in rows:
            cells = [cell.strip() for cell in row]
            if any(cells):
                extracted = " | ".join(cells)
                extracted_characters = _next_extracted_size(
                    extracted_characters,
                    extracted,
                    has_previous=bool(extracted_rows),
                    max_extracted_characters=max_extracted_characters,
                )
                extracted_rows.append(extracted)
        return "\n".join(extracted_rows)
    if suffix == ".docx":
        validated_content = preflight_zip_package(content)
        document = Document(BytesIO(validated_content))
        blocks: list[str] = []
        extracted_characters = 0
        work_units = 0
        for block in document.iter_inner_content():
            if isinstance(block, Paragraph):
                work_units += 1
                if work_units > MAX_DOCX_BLOCKS:
                    raise PreviewError("source_work_limit_exceeded")
                extracted = block.text.strip()
                if extracted:
                    extracted_characters = _next_extracted_size(
                        extracted_characters,
                        extracted,
                        has_previous=bool(blocks),
                        max_extracted_characters=max_extracted_characters,
                        separator_size=2,
                    )
                    blocks.append(extracted)
            elif isinstance(block, Table):
                for row in block.rows:
                    row_work_units = 1 + len(row.cells)
                    if work_units + row_work_units > MAX_DOCX_BLOCKS:
                        raise PreviewError("source_work_limit_exceeded")
                    work_units += row_work_units
                    extracted = " | ".join(cell.text.strip() for cell in row.cells)
                    if extracted.strip():
                        extracted_characters = _next_extracted_size(
                            extracted_characters,
                            extracted,
                            has_previous=bool(blocks),
                            max_extracted_characters=max_extracted_characters,
                            separator_size=2,
                        )
                        blocks.append(extracted)
        return "\n\n".join(blocks)
    reader = PdfReader(BytesIO(content))
    if reader.is_encrypted:
        raise PreviewError("unreadable_file")
    if len(reader.pages) > MAX_PDF_PAGES:
        raise PreviewError("source_work_limit_exceeded")
    pages: list[str] = []
    extracted_characters = 0
    for page in reader.pages:
        extracted = (page.extract_text() or "").strip()
        if extracted:
            extracted_characters = _next_extracted_size(
                extracted_characters,
                extracted,
                has_previous=bool(pages),
                max_extracted_characters=max_extracted_characters,
                separator_size=2,
            )
            pages.append(extracted)
    return "\n\n".join(pages)


def safe_source_name(filename: str) -> str:
    source_name = filename.replace("\\", "/").rsplit("/", maxsplit=1)[-1]
    if not 1 <= len(source_name) <= 255:
        raise PreviewError("validation_error")
    return source_name


def extract_office_records(
    suffix: str,
    content: bytes,
    *,
    max_extracted_characters: int,
) -> list[tuple[str, PptxLocator | XlsxLocator]]:
    validated_content = preflight_zip_package(content)
    if suffix == ".pptx":
        return _extract_pptx_records(validated_content, max_extracted_characters)
    if suffix == ".xlsx":
        return _extract_xlsx_records(validated_content, max_extracted_characters)
    raise PreviewError("unsupported_file_type")


def _extract_pptx_records(
    content: bytes, max_extracted_characters: int
) -> list[tuple[str, PptxLocator | XlsxLocator]]:
    presentation = Presentation(BytesIO(content))
    if len(presentation.slides) > MAX_PPTX_SLIDES:
        raise PreviewError("source_work_limit_exceeded")
    records: list[tuple[str, PptxLocator | XlsxLocator]] = []
    extracted_characters = 0
    shape_count = 0
    for slide_number, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        texts, title, shape_count = _pptx_shape_texts(
            slide.shapes,
            shape_count=shape_count,
            title_shape_id=title_shape.shape_id if title_shape is not None else None,
        )
        parts = ([title] if title else []) + texts
        record = "\n".join(parts)
        if not record:
            continue
        extracted_characters = _next_extracted_size(
            extracted_characters,
            record,
            has_previous=bool(records),
            max_extracted_characters=max_extracted_characters,
            separator_size=2,
        )
        locator_title = title[:MAX_PPTX_LOCATOR_TITLE_CHARS].rstrip() or None
        records.append(
            (
                record,
                PptxLocator(slide=slide_number, title=locator_title),
            )
        )
    return records


def _pptx_shape_texts(
    shapes: Collection[BaseShape],
    *,
    shape_count: int,
    title_shape_id: int | None = None,
) -> tuple[list[str], str, int]:
    shape_count += len(shapes)
    if shape_count > MAX_PPTX_SHAPES:
        raise PreviewError("source_work_limit_exceeded")
    texts: list[str] = []
    title = ""
    ordered_shapes = sorted(
        enumerate(shapes),
        key=lambda item: (item[1].top, item[1].left, item[0]),
    )
    for _, shape in ordered_shapes:
        if getattr(shape, "has_text_frame", False):
            text = shape.text.strip()
            if shape.shape_id == title_shape_id:
                title = text
            elif text:
                texts.append(text)
        if getattr(shape, "has_table", False):
            rows = [row.cells for row in shape.table.rows]
            shape_count += sum(len(cells) for cells in rows)
            if shape_count > MAX_PPTX_SHAPES:
                raise PreviewError("source_work_limit_exceeded")
            for cells in rows:
                for cell in cells:
                    text = cell.text.strip()
                    if text:
                        texts.append(text)
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            nested_texts, _, shape_count = _pptx_shape_texts(
                shape.shapes,
                shape_count=shape_count,
            )
            texts.extend(nested_texts)
    return texts, title, shape_count


def _extract_xlsx_records(
    content: bytes, max_extracted_characters: int
) -> list[tuple[str, PptxLocator | XlsxLocator]]:
    _reject_unsafe_workbook_features(content)
    workbook = load_workbook(
        BytesIO(content),
        read_only=True,
        data_only=True,
        keep_links=False,
    )
    try:
        if len(workbook.worksheets) > MAX_XLSX_SHEETS:
            raise PreviewError("source_work_limit_exceeded")
        records: list[tuple[str, PptxLocator | XlsxLocator]] = []
        extracted_characters = 0
        visited_rows = 0
        visited_cells = 0
        for sheet in workbook.worksheets:
            if sheet.max_row > MAX_XLSX_ROWS or (
                sheet.max_row * sheet.max_column > MAX_XLSX_CELLS
            ):
                raise PreviewError("source_work_limit_exceeded")
            lines: list[str] = []
            min_row = min_column = None
            max_row = max_column = 0
            for row in sheet.iter_rows():
                visited_rows += 1
                visited_cells += len(row)
                if visited_rows > MAX_XLSX_ROWS or visited_cells > MAX_XLSX_CELLS:
                    raise PreviewError("source_work_limit_exceeded")
                values: list[str] = []
                for cell in row:
                    value = _xlsx_cell_text(cell.value)
                    if value is None:
                        continue
                    values.append(value)
                    min_row = cell.row if min_row is None else min(min_row, cell.row)
                    min_column = (
                        cell.column
                        if min_column is None
                        else min(min_column, cell.column)
                    )
                    max_row = max(max_row, cell.row)
                    max_column = max(max_column, cell.column)
                if values:
                    lines.append(" | ".join(values))
            if not lines or min_row is None or min_column is None:
                continue
            record = "\n".join(lines)
            extracted_characters = _next_extracted_size(
                extracted_characters,
                record,
                has_previous=bool(records),
                max_extracted_characters=max_extracted_characters,
                separator_size=2,
            )
            cell_range = (
                f"{get_column_letter(min_column)}{min_row}:"
                f"{get_column_letter(max_column)}{max_row}"
            )
            records.append(
                (record, XlsxLocator(sheet=sheet.title, cell_range=cell_range))
            )
        return records
    finally:
        workbook.close()


def _xlsx_cell_text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    text = str(value).strip()
    return text or None


def _reject_unsafe_workbook_features(content: bytes) -> None:
    with ZipFile(BytesIO(content)) as archive:
        names = [name.casefold() for name in archive.namelist()]
        if any(
            name.endswith("vbaproject.bin") or name.startswith("xl/externallinks/")
            for name in names
        ):
            raise PreviewError("unreadable_file")
        try:
            content_types = archive.read("[Content_Types].xml").lower()
        except KeyError:
            raise PreviewError("unreadable_file") from None
    if b"macroenabled" in content_types or b"externallink" in content_types:
        raise PreviewError("unreadable_file")


def image_metadata(suffix: str, content: bytes) -> str:
    expected_format = {
        ".png": "PNG",
        ".jpg": "JPEG",
        ".jpeg": "JPEG",
        ".webp": "WEBP",
    }[suffix]
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(BytesIO(content)) as image:
                image.verify()
            with Image.open(BytesIO(content)) as image:
                actual_format = image.format
                width, height = image.size
    except (Image.DecompressionBombError, Image.DecompressionBombWarning):
        raise PreviewError("source_work_limit_exceeded") from None
    except (UnidentifiedImageError, OSError, SyntaxError):
        raise PreviewError("unreadable_file") from None
    if actual_format != expected_format:
        raise PreviewError("unreadable_file")
    if (
        width <= 0
        or height <= 0
        or width > MAX_IMAGE_DIMENSION
        or height > MAX_IMAGE_DIMENSION
        or width * height > MAX_IMAGE_PIXELS
    ):
        raise PreviewError("source_work_limit_exceeded")
    return f"{actual_format} image, {width} x {height} pixels"


def _preview_records(
    source_name: str,
    records: list[tuple[str, SourceLocator]],
    settings: ChunkSettings,
) -> PreviewResponse:
    if not records:
        raise PreviewError("empty_source")
    visible_records = records[:MAX_PREVIEW_ITEMS]
    preview_warnings = []
    if settings.training_type == "qa":
        preview_warnings.append(PreviewWarning(code="qa_generated_after_processing"))
    if len(records) > MAX_PREVIEW_ITEMS:
        preview_warnings.append(PreviewWarning(code="preview_truncated"))
    return PreviewResponse(
        source_name=source_name,
        character_count=sum(len(text) for text, _ in records),
        items=[
            PreviewItem(
                position=position,
                text=text[:MAX_PREVIEW_EXCERPT_CHARS],
                character_count=len(text),
                truncated=len(text) > MAX_PREVIEW_EXCERPT_CHARS,
                locator=locator,
            )
            for position, (text, locator) in enumerate(visible_records, start=1)
        ],
        warnings=preview_warnings,
    )


def preflight_zip_package(content: bytes) -> bytes:
    validated_content = _raw_zip_preflight(content)
    with ZipFile(BytesIO(validated_content)) as archive:
        members = archive.infolist()
    if len(members) > MAX_ZIP_ARCHIVE_MEMBERS:
        raise PreviewError("unsafe_archive")
    uncompressed_bytes = sum(member.file_size for member in members)
    compressed_bytes = sum(member.compress_size for member in members)
    if uncompressed_bytes > MAX_ZIP_UNCOMPRESSED_BYTES:
        raise PreviewError("unsafe_archive")
    if uncompressed_bytes and compressed_bytes == 0:
        raise PreviewError("unsafe_archive")
    if compressed_bytes and uncompressed_bytes / compressed_bytes > MAX_ZIP_COMPRESSION_RATIO:
        raise PreviewError("unsafe_archive")
    if any(
        member.file_size
        and (
            member.compress_size == 0
            or member.file_size / member.compress_size > MAX_ZIP_COMPRESSION_RATIO
        )
        for member in members
    ):
        raise PreviewError("unsafe_archive")
    return validated_content


def _raw_zip_preflight(content: bytes) -> bytes:
    """Validate bounded ZIP metadata before ZipFile can allocate ZipInfo objects."""

    eocd = _find_valid_eocd(content)
    (
        eocd_offset,
        entry_count,
        central_directory_size,
        central_directory_offset,
        comment_length,
    ) = eocd
    if entry_count == 0 or entry_count > MAX_ZIP_ARCHIVE_MEMBERS:
        raise PreviewError("unsafe_archive")
    if central_directory_size > MAX_ZIP_CENTRAL_DIRECTORY_BYTES:
        raise PreviewError("unsafe_archive")
    central_directory_end = central_directory_offset + central_directory_size
    if (
        central_directory_offset > eocd_offset
        or central_directory_end != eocd_offset
    ):
        raise PreviewError("unsafe_archive")

    _validate_raw_central_directory(
        content,
        entry_count=entry_count,
        start=central_directory_offset,
        end=central_directory_end,
    )
    if comment_length == 0:
        return content

    # CPython's zipfile searches for the last EOCD signature, including inside
    # the comment. Remove an already-validated comment before invoking it.
    sanitized = bytearray(content[: eocd_offset + ZIP_EOCD_SIZE])
    struct.pack_into("<H", sanitized, eocd_offset + 20, 0)
    return bytes(sanitized)


def _find_valid_eocd(content: bytes) -> tuple[int, int, int, int, int]:
    search_start = max(0, len(content) - ZIP_EOCD_SIZE - ZIP_MAX_COMMENT_BYTES)
    candidate = content.rfind(ZIP_EOCD_SIGNATURE, search_start)
    if candidate < 0 and content.find(ZIP_EOCD_SIGNATURE) < 0:
        raise PreviewError("unreadable_file")
    while candidate >= search_start:
        parsed = _parse_eocd_candidate(content, candidate)
        if parsed is not None:
            return parsed
        candidate = content.rfind(ZIP_EOCD_SIGNATURE, search_start, candidate)
    raise PreviewError("unsafe_archive")


def _parse_eocd_candidate(
    content: bytes, offset: int
) -> tuple[int, int, int, int, int] | None:
    if offset + ZIP_EOCD_SIZE > len(content):
        return None
    (
        signature,
        disk_number,
        central_directory_disk,
        entries_on_disk,
        entry_count,
        central_directory_size,
        central_directory_offset,
        comment_length,
    ) = struct.unpack_from("<4s4H2IH", content, offset)
    if signature != ZIP_EOCD_SIGNATURE:
        return None
    if offset + ZIP_EOCD_SIZE + comment_length != len(content):
        return None
    if (
        disk_number != 0
        or central_directory_disk != 0
        or entries_on_disk != entry_count
        or entry_count == ZIP16_SENTINEL
        or central_directory_size == ZIP32_SENTINEL
        or central_directory_offset == ZIP32_SENTINEL
    ):
        return None
    if (
        offset >= 20
        and content[offset - 20 : offset - 16] == ZIP64_LOCATOR_SIGNATURE
    ):
        return None
    if (
        content.find(
            ZIP64_EOCD_SIGNATURE,
            max(0, central_directory_offset - 56),
            offset,
        )
        >= 0
    ):
        return None
    if central_directory_offset + central_directory_size != offset:
        return None
    return (
        offset,
        entry_count,
        central_directory_size,
        central_directory_offset,
        comment_length,
    )


def _validate_raw_central_directory(
    content: bytes,
    *,
    entry_count: int,
    start: int,
    end: int,
) -> None:
    position = start
    parsed_entries = 0
    total_compressed_bytes = 0
    total_uncompressed_bytes = 0
    while position < end:
        if (
            parsed_entries >= MAX_ZIP_ARCHIVE_MEMBERS
            or position + ZIP_CENTRAL_DIRECTORY_HEADER_SIZE > end
            or content[position : position + 4] != ZIP_CENTRAL_DIRECTORY_SIGNATURE
        ):
            raise PreviewError("unsafe_archive")
        compressed_size, uncompressed_size = struct.unpack_from(
            "<II", content, position + 20
        )
        filename_length, extra_length, member_comment_length = struct.unpack_from(
            "<HHH", content, position + 28
        )
        member_disk = struct.unpack_from("<H", content, position + 34)[0]
        local_header_offset = struct.unpack_from("<I", content, position + 42)[0]
        if (
            compressed_size == ZIP32_SENTINEL
            or uncompressed_size == ZIP32_SENTINEL
            or local_header_offset == ZIP32_SENTINEL
            or member_disk != 0
        ):
            raise PreviewError("unsafe_archive")

        record_end = (
            position
            + ZIP_CENTRAL_DIRECTORY_HEADER_SIZE
            + filename_length
            + extra_length
            + member_comment_length
        )
        if record_end > end or local_header_offset + 4 > start:
            raise PreviewError("unsafe_archive")
        if content[local_header_offset : local_header_offset + 4] != ZIP_LOCAL_FILE_SIGNATURE:
            raise PreviewError("unsafe_archive")

        extra_start = position + ZIP_CENTRAL_DIRECTORY_HEADER_SIZE + filename_length
        _validate_zip_extra_fields(content, start=extra_start, end=extra_start + extra_length)

        total_compressed_bytes += compressed_size
        total_uncompressed_bytes += uncompressed_size
        if total_uncompressed_bytes > MAX_ZIP_UNCOMPRESSED_BYTES:
            raise PreviewError("unsafe_archive")
        if uncompressed_size and (
            compressed_size == 0
            or uncompressed_size / compressed_size > MAX_ZIP_COMPRESSION_RATIO
        ):
            raise PreviewError("unsafe_archive")

        parsed_entries += 1
        position = record_end

    if parsed_entries != entry_count or position != end:
        raise PreviewError("unsafe_archive")
    if total_uncompressed_bytes and total_compressed_bytes == 0:
        raise PreviewError("unsafe_archive")
    if (
        total_compressed_bytes
        and total_uncompressed_bytes / total_compressed_bytes
        > MAX_ZIP_COMPRESSION_RATIO
    ):
        raise PreviewError("unsafe_archive")


def _validate_zip_extra_fields(content: bytes, *, start: int, end: int) -> None:
    position = start
    while position < end:
        if position + 4 > end:
            raise PreviewError("unsafe_archive")
        field_id, field_size = struct.unpack_from("<HH", content, position)
        position += 4
        if position + field_size > end or field_id == ZIP64_EXTRA_FIELD_ID:
            raise PreviewError("unsafe_archive")
        position += field_size


def _next_extracted_size(
    current_size: int,
    text: str,
    *,
    has_previous: bool,
    max_extracted_characters: int,
    separator_size: int = 1,
) -> int:
    next_size = current_size + len(text) + (separator_size if has_previous else 0)
    if next_size > max_extracted_characters:
        raise PreviewError("source_too_large")
    return next_size


def _bounded_text(text: str, max_extracted_characters: int) -> str:
    if len(text) > max_extracted_characters:
        raise PreviewError("source_too_large")
    return text


def _validate_file_size(content: bytes, max_upload_bytes: int) -> None:
    _require_positive_limit(max_upload_bytes)
    if len(content) > max_upload_bytes:
        raise PreviewError("file_too_large")


def _validate_text_budget(
    text: str,
    *,
    max_text_bytes: int | None,
    max_characters: int | None,
) -> None:
    if max_text_bytes is not None:
        _require_positive_limit(max_text_bytes)
        try:
            text_bytes = len(text.encode("utf-8"))
        except UnicodeEncodeError:
            raise PreviewError("unreadable_file") from None
        if text_bytes > max_text_bytes:
            raise PreviewError("text_too_large")
    if max_characters is not None:
        _require_positive_limit(max_characters)
        if len(text) > max_characters:
            raise PreviewError("source_too_large")


def _require_positive_limit(limit: int) -> None:
    if type(limit) is not int or limit <= 0:
        raise ValueError("preview resource limits must be positive integers")


def _segments(text: str, settings: ChunkSettings) -> Iterator[str]:
    if settings.setting_mode == "auto":
        settings = ChunkSettings()
    split_mode = settings.split_mode
    if split_mode == "size":
        yield from _split_to_size(text, settings.chunk_size)
        return
    if split_mode == "char":
        candidates = _split_on_delimiter(text, settings.splitter)
    else:
        candidates = _split_paragraphs(text)
    for candidate in candidates:
        candidate = candidate.strip()
        if candidate:
            yield from _split_to_size(candidate, settings.chunk_size)


def _split_on_delimiter(text: str, delimiter: str) -> Iterator[str]:
    start = 0
    while True:
        boundary = text.find(delimiter, start)
        if boundary < 0:
            yield text[start:]
            return
        yield text[start:boundary]
        start = boundary + len(delimiter)


def _split_paragraphs(text: str) -> Iterator[str]:
    start = 0
    for boundary in re.finditer(r"(?:\r?\n[ \t]*){2,}", text):
        yield text[start : boundary.start()]
        start = boundary.end()
    yield text[start:]


def _split_to_size(text: str, size: int) -> Iterator[str]:
    for start in range(0, len(text), size):
        segment = text[start : start + size]
        if segment:
            yield segment
